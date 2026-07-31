"""
fiche_writer.py  (v2)
=====================
Fill a complete BIA fiche (.docx) template from form_data collected by the
web editor.  All keys in form_data are optional.

form_data schema:
  {
    "entity": {
      "nom_responsable": str,
      "organisation":    str,
      "contraintes":     str,
      "periodes_critiques": str,
      "historique_interruptions": str
    },
    "participants": [{"nom": str, "fonction": str}, …],
    "activities": [
      {
        "name":               str,
        "description":        str,
        "ressources_utilisees": str,
        "periode_critique":   str,
        "criticite":          str,
        "dmia_exprimee":      str,
        "premieres_actions":  str,
        "impacts": {
          "Image de marque":            {"A": str, "B": str},
          "Désorganisation interne":    {"A": str, "B": str},
          "Juridique et réglementaire": {"A": str, "B": str},
          "Financier":                  {"A": str, "B": str}
        }
      }
    ],
    "echanges": [
      {"groupes": str, "ie": str, "type_info": str, "tr": str, "ressources_si": str}
    ],
    "montee_en_charge": {
      "Effectif":           {"Nominal": str, "H0": str, …},
      "Positions":          {…},
      "Télétravail":        {…},
      "Effectif cumulé":    {…},
      "Positions cumulées": {…},
      "Télétravail cumulé": {…},
      "% Effectif cumulé":  {…}
    },
    "collaborateurs_cles": [{"fonction": str, "nom": str, "prenom": str, "suppleants": str}],
    "applications":         [{"application": str, "criticite": str, "dmia": str, "pmdt": str, "commentaires": str}],
    "app_availability":     [{"designation": str, "H+2": str, "H+4": str, "J+1": str, …}],
    "documents":            [{"document": str, "stockage": str, "duplication": str}],
    "observations":         str
  }
"""
from __future__ import annotations

import re
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement

_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

# ─── constants ────────────────────────────────────────────────────────────────

TIME_COLS: list[str] = [
    "Nominal", "H0", "H+2", "H+4",
    "J+1", "J+2", "J+3", "J+5", "J+10", "J+15",
]
TIME_COL_MINUTES: dict[str, int] = {
    "Nominal": -1,
    "H0": 0, "H+2": 120, "H+4": 240,
    "J+1": 1440, "J+2": 2880, "J+3": 4320,
    "J+5": 7200, "J+10": 14400, "J+15": 21600,
}
METRIC_ROWS: list[str] = [
    "Effectif", "Positions", "Télétravail",
    "Effectif cumulé", "Positions cumulées",
    "Télétravail cumulé", "% Effectif cumulé",
]
IMPACT_LABELS: list[str] = [
    "Image de marque",
    "Désorganisation interne",
    "Juridique et réglementaire",
    "Financier",
]
# Row-labels in the Fiche d'identité table and their form_data keys
IDENTITY_ROWS: list[tuple[str, str]] = [
    ("responsable",     "nom_responsable"),
    ("organisation",    "organisation"),
    ("présentation",    None),           # auto from activities → skip manual fill here
    ("contrainte",      "contraintes"),
    ("période",         "periodes_critiques"),
    ("historique",      "historique_interruptions"),
]


def parse_dmia_minutes(dmia: str) -> int:
    """
    Parse a DMIA string into minutes.
    Handles all formats found across Devoteam synthèse files:
      H+4, 4H, H4  |  J+1, 1J, J1  |  1 Jour, 2 Jours  |  1 Heure, 4 Heures
      30 minutes  |  1 semaine, 2 semaines  |  H0 / Immédiat = 0

    Returns:
        >= 0   valid duration in minutes  (0 = immediate / H0)
        -1     unrecognised or non-time string (pure number, code, name, etc.)
    """
    if not dmia or not dmia.strip():
        return -1
    s     = re.sub(r"\s+", " ", dmia.strip().upper())
    s_nsp = s.replace(" ", "")

    # ── immediate ─────────────────────────────────────────────────────────────
    if s_nsp in ("H0", "H+0", "IMMEDIAT", "IMMÉDIAT", "IMMEDIAT"):
        return 0

    # ── compact: H+4 / H4 / 4H ────────────────────────────────────────────────
    m = re.match(r"^H\+?(\d+(?:[.,]\d+)?)$", s_nsp)
    if m:
        return round(float(m.group(1).replace(",", ".")) * 60)
    m = re.match(r"^(\d+(?:[.,]\d+)?)H$", s_nsp)
    if m:
        return round(float(m.group(1).replace(",", ".")) * 60)

    # ── compact: J+1 / J1 / 1J ────────────────────────────────────────────────
    m = re.match(r"^J\+?(\d+(?:[.,]\d+)?)$", s_nsp)
    if m:
        return round(float(m.group(1).replace(",", ".")) * 1440)
    m = re.match(r"^(\d+(?:[.,]\d+)?)J$", s_nsp)
    if m:
        return round(float(m.group(1).replace(",", ".")) * 1440)

    # ── French word forms ──────────────────────────────────────────────────────
    m = re.match(r"^(\d+(?:[.,]\d+)?)\s+JOUR[S]?$", s)
    if m:
        return round(float(m.group(1).replace(",", ".")) * 1440)
    m = re.match(r"^(\d+(?:[.,]\d+)?)\s+HEURE[S]?$", s)
    if m:
        return round(float(m.group(1).replace(",", ".")) * 60)
    m = re.match(r"^(\d+(?:[.,]\d+)?)\s+(?:MINUTE[S]?|MIN)$", s)
    if m:
        return round(float(m.group(1).replace(",", ".")))
    m = re.match(r"^(\d+(?:[.,]\d+)?)\s+SEMAINE[S]?$", s)
    if m:
        return round(float(m.group(1).replace(",", ".")) * 10080)
    m = re.match(r"^(\d+(?:[.,]\d+)?)\s+MOIS$", s)
    if m:
        return round(float(m.group(1).replace(",", ".")) * 43200)

    # ── mixed: "2H30", "1J12H", "1 J 4 H" ────────────────────────────────────
    m = re.match(r"^(\d+)H(\d+)(?:MIN)?$", s_nsp)
    if m:
        return int(m.group(1)) * 60 + int(m.group(2))
    m = re.match(r"^(\d+)J(\d+)H$", s_nsp)
    if m:
        return int(m.group(1)) * 1440 + int(m.group(2)) * 60

    # ── not a recognised duration ──────────────────────────────────────────────
    return -1


# ─── low-level XML helpers ────────────────────────────────────────────────────

def _elem_text(elem) -> str:
    return "".join((t.text or "") for t in elem.iter(f"{{{_W}}}t")).strip()

def _cell_text(cell) -> str:
    return cell.text.strip()

def _set_para_text(para_elem, text: str) -> None:
    for r in list(para_elem.findall(f"{{{_W}}}r")):
        para_elem.remove(r)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    if text and (text[0] == " " or text[-1] == " "):
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    r.append(t)
    para_elem.append(r)

def _clear_data_rows(table, header_rows: int = 1) -> None:
    tbl = table._tbl
    for row in list(tbl.findall(f"{{{_W}}}tr"))[header_rows:]:
        tbl.remove(row)

def _set_cell_text(cell, text: str) -> None:
    tc = cell._tc
    paras = tc.findall(f".//{{{_W}}}p")
    if paras:
        _set_para_text(paras[0], str(text))
        for extra in paras[1:]:
            tc.remove(extra)

def _clone_row_with_values(template_tr, values: list[str]):
    new_tr = deepcopy(template_tr)
    tcs = new_tr.findall(f"{{{_W}}}tc")
    for tc, val in zip(tcs, values):
        paras = tc.findall(f".//{{{_W}}}p")
        if paras:
            _set_para_text(paras[0], str(val))
            for extra in paras[1:]:
                tc.remove(extra)
    return new_tr


# ─── public API ───────────────────────────────────────────────────────────────

def fill_fiche(template_path: Path, form_data: dict, output_path: Path) -> None:
    doc = Document(str(template_path))
    activities   = form_data.get("activities", [])
    entity       = form_data.get("entity", {})
    participants = form_data.get("participants", [])
    echanges     = form_data.get("echanges", [])
    montee       = form_data.get("montee_en_charge", {})
    collabs      = form_data.get("collaborateurs_cles", [])
    apps         = form_data.get("applications", [])
    app_avail    = form_data.get("app_availability", [])
    docs         = form_data.get("documents", [])
    observations = form_data.get("observations", "")

    # Participants (Fiche de suivi)
    if participants:
        _fill_participants(doc, participants)

    # Fiche d'identité (Table 7) — includes §3 presentation
    _fill_entity_identity(doc, entity, activities)

    if activities:
        # §4 Description et criticité
        _fill_section4_criticite_full(doc, activities)
        # §5.2 Évaluation des impacts
        _fill_section52_impacts_full(doc, activities)
        # §5.3 DMIA
        _fill_section53_dmia(doc, activities)

    # §6 Analyse des échanges
    if echanges:
        _fill_exchanges(doc, echanges)

    # §7.1 Montée en charge — 7 shared rows
    _fill_section71_montee_v2(doc, montee, activities)

    # §7.2 Collaborateurs clés
    if collabs:
        _fill_collaborateurs(doc, collabs)

    # §7.3 Applications + availability
    if apps:
        _fill_applications(doc, apps)
    if app_avail:
        _fill_app_availability(doc, app_avail)

    # §7.4 Documents et fichiers
    if docs:
        _fill_documents(doc, docs)

    # Observations
    if observations:
        _fill_observations(doc, observations)

    doc.save(str(output_path))


# ─── section fillers ──────────────────────────────────────────────────────────

def _fill_participants(doc: Document, participants: list[dict]) -> None:
    """Fiche de suivi – participant rows."""
    for table in doc.tables:
        if not table.rows:
            continue
        h0 = _cell_text(table.rows[0].cells[0]).lower()
        if "entit" not in h0:
            continue
        if len(table.rows) < 2:
            continue
        h1 = _cell_text(table.rows[1].cells[0]).lower()
        if "sent" not in h1 and "pr" not in h1:
            continue
        # rows 0 = entity header, row 1 = column header; data rows 2+
        template_tr = (
            deepcopy(table.rows[2]._tr) if len(table.rows) > 2 else None
        )
        _clear_data_rows(table, header_rows=2)
        if template_tr is not None:
            for p in participants:
                vals = ["Présents", p.get("nom", ""), p.get("fonction", "")]
                table._tbl.append(_clone_row_with_values(template_tr, vals))
        return


def _fill_entity_identity(doc: Document, entity: dict, activities: list[dict]) -> None:
    """Fiche d'identité (6-row key-value table) — also handles §3 presentation."""
    for table in doc.tables:
        if len(table.rows) < 2:
            continue
        h0 = _cell_text(table.rows[0].cells[0]).lower()
        h1 = _cell_text(table.rows[1].cells[0]).lower()
        if "responsable" not in h0 or "organisation" not in h1:
            continue
        # Found the identity table
        for row in table.rows:
            if len(row.cells) < 2:
                continue
            label = _cell_text(row.cells[0]).lower()
            val_cell = row.cells[1]

            if "responsable" in label:
                if entity.get("nom_responsable"):
                    _set_cell_text(val_cell, entity["nom_responsable"])

            elif "organisation" in label:
                if entity.get("organisation"):
                    _set_cell_text(val_cell, entity["organisation"])

            elif "pr" in label and "sentation" in label:
                # §3 – auto-build from activities
                if activities:
                    lines = (
                        ["Les activités principales de l'entité sont les suivantes :"]
                        + [f"- {a['name']} : {a.get('description', '')}"
                           for a in activities]
                    )
                    tc = val_cell._tc
                    paras = tc.findall(f"{{{_W}}}p")
                    if paras:
                        _set_para_text(paras[0], lines[0])
                        for extra in paras[1:]:
                            tc.remove(extra)
                        for line in lines[1:]:
                            new_p = deepcopy(paras[0])
                            _set_para_text(new_p, line)
                            tc.append(new_p)

            elif "contrainte" in label:
                if entity.get("contraintes"):
                    _set_cell_text(val_cell, entity["contraintes"])

            elif "p" in label and "riode" in label and "crit" in label:
                if entity.get("periodes_critiques"):
                    _set_cell_text(val_cell, entity["periodes_critiques"])

            elif "historique" in label or "interruption" in label:
                if entity.get("historique_interruptions"):
                    _set_cell_text(val_cell, entity["historique_interruptions"])
        return


def _fill_section4_criticite_full(doc: Document, activities: list[dict]) -> None:
    """§4 Description et criticité – all 4 columns per activity."""
    for table in doc.tables:
        if not table.rows:
            continue
        combined = " ".join(_cell_text(c) for c in table.rows[0].cells).lower()
        if "activit" in combined and "criticit" in combined:
            template_tr = (
                deepcopy(table.rows[1]._tr) if len(table.rows) > 1 else None
            )
            _clear_data_rows(table, header_rows=1)
            if template_tr is not None:
                n_cols = len(table.rows[0].cells)
                for act in activities:
                    vals = [
                        act.get("name", ""),
                        act.get("ressources_utilisees", ""),
                        act.get("periode_critique", ""),
                        act.get("criticite", ""),
                    ]
                    vals = (vals + [""] * n_cols)[:n_cols]
                    table._tbl.append(_clone_row_with_values(template_tr, vals))
            return


def _fill_section52_impacts_full(doc: Document, activities: list[dict]) -> None:
    """§5.2 Évaluation des impacts – clone template table, fill scores."""
    body = doc.element.body
    body_children = list(body)

    # Locate the evaluation table (no "poids", has "image de marque" in row1)
    impact_tbl = None
    impact_idx = None
    for i, child in enumerate(body_children):
        if child.tag != f"{{{_W}}}tbl":
            continue
        rows = child.findall(f"{{{_W}}}tr")
        if len(rows) < 2:
            continue
        row0_tcs = rows[0].findall(f"{{{_W}}}tc")
        row1_tcs = rows[1].findall(f"{{{_W}}}tc")
        if not row0_tcs or not row1_tcs:
            continue
        row1_c0 = _elem_text(row1_tcs[0]).lower()
        row0_all = " ".join(_elem_text(c) for c in row0_tcs).lower()
        if "image de marque" in row1_c0 and "poids" not in row0_all:
            impact_tbl = child
            impact_idx = i
            break

    if impact_tbl is None:
        return

    caption_elem = None
    if impact_idx and impact_idx > 0:
        prev = body_children[impact_idx - 1]
        if prev.tag == f"{{{_W}}}p":
            txt = _elem_text(prev).lower()
            if "tableau" in txt or "xx" in txt:
                caption_elem = prev

    def _fill_one_impact_table(tbl_elem, act: dict) -> None:
        """Set activity name in row0 and fill score values in rows 1-4."""
        rows = tbl_elem.findall(f"{{{_W}}}tr")
        if not rows:
            return
        # Row 0 – activity name in col 0
        cells0 = rows[0].findall(f"{{{_W}}}tc")
        if cells0:
            paras = cells0[0].findall(f".//{{{_W}}}p")
            if paras:
                _set_para_text(paras[0], act.get("name", ""))
        # Rows 1–4 – score values
        impacts = act.get("impacts", {})
        for ri, row_elem in enumerate(rows[1:5], start=1):
            tcs = row_elem.findall(f"{{{_W}}}tc")
            if len(tcs) < 3:
                continue
            label = _elem_text(tcs[0])
            # match by row index fallback
            impact_key = label if label in impacts else (
                IMPACT_LABELS[ri - 1] if ri - 1 < len(IMPACT_LABELS) else None
            )
            if impact_key and impact_key in impacts:
                scores = impacts[impact_key]
                for tc, col_key in zip(tcs[1:3], ("A", "B")):
                    paras = tc.findall(f".//{{{_W}}}p")
                    if paras:
                        _set_para_text(paras[0], str(scores.get(col_key, "")))

    _fill_one_impact_table(impact_tbl, activities[0])
    if caption_elem is not None:
        _replace_xx_in_para(caption_elem, activities[0].get("name", ""))

    insert_after = impact_tbl
    for act in activities[1:]:
        new_tbl = deepcopy(impact_tbl)
        _fill_one_impact_table(new_tbl, act)
        if caption_elem is not None:
            new_cap = deepcopy(caption_elem)
            _replace_xx_in_para(new_cap, act.get("name", ""))
            insert_after.addnext(new_tbl)
            insert_after.addnext(new_cap)
        else:
            insert_after.addnext(new_tbl)
        insert_after = new_tbl


def _fill_section53_dmia(doc: Document, activities: list[dict]) -> None:
    """§5.3 DMIA – Processus | DMIA Exprimée | Premières actions."""
    for table in doc.tables:
        if not table.rows:
            continue
        combined = " ".join(_cell_text(c) for c in table.rows[0].cells).lower()
        if "processus" in combined and "dmia" in combined:
            template_tr = (
                deepcopy(table.rows[1]._tr) if len(table.rows) > 1 else None
            )
            _clear_data_rows(table, header_rows=1)
            if template_tr is not None:
                n_cols = len(table.rows[0].cells)
                for act in activities:
                    vals = [
                        act.get("name", ""),
                        act.get("dmia_exprimee", ""),
                        act.get("premieres_actions", ""),
                    ]
                    vals = (vals + [""] * n_cols)[:n_cols]
                    table._tbl.append(_clone_row_with_values(template_tr, vals))
            return


def _fill_exchanges(doc: Document, echanges: list[dict]) -> None:
    """§6 Analyse des échanges d'information."""
    for table in doc.tables:
        if not table.rows:
            continue
        combined = " ".join(_cell_text(c) for c in table.rows[0].cells).lower()
        if "groupes fonctionnels" in combined or (
                "interne" in combined and "externe" in combined
                and "ressources si" in combined):
            template_tr = (
                deepcopy(table.rows[1]._tr) if len(table.rows) > 1 else None
            )
            _clear_data_rows(table, header_rows=1)
            if template_tr is not None:
                n_cols = len(table.rows[0].cells)
                for ech in echanges:
                    vals = [
                        ech.get("groupes", ""),
                        ech.get("ie", ""),
                        ech.get("type_info", ""),
                        ech.get("tr", ""),
                        ech.get("ressources_si", ""),
                    ]
                    vals = (vals + [""] * n_cols)[:n_cols]
                    table._tbl.append(_clone_row_with_values(template_tr, vals))
            return


def _fill_section71_montee_v2(
    doc: Document,
    montee: dict,
    activities: list[dict],
) -> None:
    """
    §7.1 Montée en charge – 7 shared metric rows (NOT duplicated per activity).
    Commentaires = all activity names + DMIA reminders, joined by newlines.
    """
    for table in doc.tables:
        if not table.rows:
            continue
        combined = " ".join(_cell_text(c) for c in table.rows[0].cells).lower()
        if ("effectif" in combined or "mont" in combined) and "commentaires" in combined:
            col_headers = [_cell_text(c) for c in table.rows[0].cells]
            col_idx_map = {h.strip(): i for i, h in enumerate(col_headers) if h.strip()}
            n_cols = len(col_headers)

            n_data = len(table.rows) - 1
            template_rows = [
                deepcopy(table.rows[i]._tr)
                for i in range(1, min(n_data + 1, len(METRIC_ROWS) + 1))
            ]
            if not template_rows:
                return

            _clear_data_rows(table, header_rows=1)

            # Commentaires text = all activities with DMIA reminder
            comment_parts = []
            for act in activities:
                part = act.get("name", "")
                dmia = act.get("dmia_exprimee", "")
                if dmia:
                    part += f" (DMIA: {dmia})"
                if part.strip():
                    comment_parts.append(part)
            comment_text = "\n".join(comment_parts)

            for row_idx, metric in enumerate(METRIC_ROWS):
                template_tr = template_rows[row_idx % len(template_rows)]
                vals = [""] * n_cols
                vals[0] = metric

                for col_name in TIME_COLS:
                    idx = col_idx_map.get(col_name)
                    if idx is not None:
                        vals[idx] = str(
                            montee.get(metric, {}).get(col_name, "") or ""
                        )

                comm_idx = col_idx_map.get("Commentaires")
                if comm_idx is not None:
                    vals[comm_idx] = comment_text

                table._tbl.append(_clone_row_with_values(template_tr, vals))
            return


def _fill_collaborateurs(doc: Document, collabs: list[dict]) -> None:
    """§7.2 Collaborateurs clés."""
    for table in doc.tables:
        if not table.rows:
            continue
        combined = " ".join(_cell_text(c) for c in table.rows[0].cells).lower()
        if "suppl" in combined and ("nom" in combined or "fonction" in combined):
            template_tr = (
                deepcopy(table.rows[1]._tr) if len(table.rows) > 1 else None
            )
            _clear_data_rows(table, header_rows=1)
            if template_tr is not None:
                n_cols = len(table.rows[0].cells)
                for c in collabs:
                    vals = [
                        c.get("fonction", ""),
                        c.get("nom", ""),
                        c.get("prenom", ""),
                        c.get("suppleants", ""),
                    ]
                    vals = (vals + [""] * n_cols)[:n_cols]
                    table._tbl.append(_clone_row_with_values(template_tr, vals))
            return


def _fill_applications(doc: Document, apps: list[dict]) -> None:
    """§7.3 Applications informatiques."""
    for table in doc.tables:
        if not table.rows:
            continue
        combined = " ".join(_cell_text(c) for c in table.rows[0].cells).lower()
        if "application" in combined and "pmdt" in combined:
            template_tr = (
                deepcopy(table.rows[1]._tr) if len(table.rows) > 1 else None
            )
            _clear_data_rows(table, header_rows=1)
            if template_tr is not None:
                n_cols = len(table.rows[0].cells)
                for a in apps:
                    vals = [
                        a.get("application", ""),
                        a.get("criticite", ""),
                        a.get("dmia", ""),
                        a.get("pmdt", ""),
                        a.get("commentaires", ""),
                    ]
                    vals = (vals + [""] * n_cols)[:n_cols]
                    table._tbl.append(_clone_row_with_values(template_tr, vals))
            return


def _fill_app_availability(doc: Document, app_avail: list[dict]) -> None:
    """§7.3b Disponibilité des applications (Désignation | H+2 | H+4 …)."""
    for table in doc.tables:
        if not table.rows:
            continue
        combined = " ".join(_cell_text(c) for c in table.rows[0].cells).lower()
        if "d" in combined and "signation" in combined and "h+2" in combined:
            col_headers = [_cell_text(c).strip() for c in table.rows[0].cells]
            col_idx_map = {h: i for i, h in enumerate(col_headers) if h}
            n_cols = len(col_headers)
            template_tr = (
                deepcopy(table.rows[1]._tr) if len(table.rows) > 1 else None
            )
            _clear_data_rows(table, header_rows=1)
            if template_tr is not None:
                avail_cols = ["H+2", "H+4", "J+1", "J+2", "J+3", "J+5", "J+10", "J+15"]
                for a in app_avail:
                    vals = [""] * n_cols
                    if col_idx_map.get("Désignation") is not None:
                        vals[col_idx_map["Désignation"]] = a.get("designation", "")
                    for col in avail_cols:
                        idx = col_idx_map.get(col)
                        if idx is not None:
                            vals[idx] = str(a.get(col, "") or "")
                    table._tbl.append(_clone_row_with_values(template_tr, vals))
            return


def _fill_documents(doc: Document, docs: list[dict]) -> None:
    """§7.4 Documents et fichiers critiques."""
    for table in doc.tables:
        if not table.rows:
            continue
        combined = " ".join(_cell_text(c) for c in table.rows[0].cells).lower()
        if "duplication" in combined and "document" in combined:
            template_tr = (
                deepcopy(table.rows[1]._tr) if len(table.rows) > 1 else None
            )
            _clear_data_rows(table, header_rows=1)
            if template_tr is not None:
                n_cols = len(table.rows[0].cells)
                for d in docs:
                    vals = [
                        d.get("document", ""),
                        d.get("stockage", ""),
                        d.get("duplication", ""),
                    ]
                    vals = (vals + [""] * n_cols)[:n_cols]
                    table._tbl.append(_clone_row_with_values(template_tr, vals))
            return


def _fill_observations(doc: Document, observations: str) -> None:
    """Observations – first paragraph after the 'Observations' heading."""
    found_heading = False
    for para in doc.paragraphs:
        txt = para.text.strip()
        if not found_heading:
            if "observation" in txt.lower() and para.style.name.startswith("Heading"):
                found_heading = True
        else:
            if txt or not para.style.name.startswith("Heading"):
                _set_para_text(para._p, observations)
                return


# ─── XML helpers ──────────────────────────────────────────────────────────────

def _set_impact_row0_name(tbl_elem, name: str) -> None:
    rows = tbl_elem.findall(f"{{{_W}}}tr")
    if not rows:
        return
    cells = rows[0].findall(f"{{{_W}}}tc")
    if not cells:
        return
    paras = cells[0].findall(f".//{{{_W}}}p")
    if paras:
        _set_para_text(paras[0], name)


def _replace_xx_in_para(para_elem, replacement: str) -> None:
    for t in para_elem.iter(f"{{{_W}}}t"):
        if t.text and re.search(r"\bxx\b", t.text, re.IGNORECASE):
            t.text = re.sub(
                r"\bxx\b", replacement, t.text, count=1, flags=re.IGNORECASE
            )
            return
