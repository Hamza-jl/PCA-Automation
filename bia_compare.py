"""
bia_compare.py — Compare two BIA synthèse Excel files.

Usage:
    result = compare_bia_files(original_path, generated_path)

Returns a dict with overall similarity, field accuracy, and per-sheet breakdowns.
"""
from __future__ import annotations

import os
from pathlib import Path
from typing import Any

import openpyxl
from rapidfuzz import fuzz, process as rfprocess

# Sheets to skip (metadata / reference only)
_SKIP_SHEETS: set[str] = {
    "Synthèse",
    "A masquer",
    "Echelle d'impact",
    "Redémarrages des applications",
}

_EMPTY_VALUES: set[str] = {"-", "nan", "none", ""}


def _cell_str(cell) -> str:
    """Convert an openpyxl cell value to a clean string."""
    v = cell.value
    if v is None:
        return ""
    s = str(v).strip()
    if s.lower() in _EMPTY_VALUES:
        return ""
    return s


def _read_sheet_data(ws) -> tuple[list[str], list[dict]]:
    """
    Read headers from row 5 and data rows from row 6+.

    Returns:
        headers  – list of column header strings (1-indexed, col 1 = index 0)
        rows     – list of dicts {col_index_1based: value_str}
                   Fully-blank rows are skipped.
    """
    headers: list[str] = []
    max_col = ws.max_column or 1

    # Row 5 → headers
    for col in range(1, max_col + 1):
        cell = ws.cell(row=5, column=col)
        headers.append(_cell_str(cell))

    # Row 6+ → data
    rows: list[dict] = []
    for row_idx in range(6, (ws.max_row or 5) + 1):
        row_dict: dict[int, str] = {}
        for col in range(1, max_col + 1):
            val = _cell_str(ws.cell(row=row_idx, column=col))
            if val:
                row_dict[col] = val
        if row_dict:  # skip fully blank rows
            rows.append(row_dict)

    return headers, rows


def _entity_key(row: dict) -> str:
    """
    Concatenate non-empty values from columns 2, 3, 4 (B, C, D) with ' / '.
    This uniquely identifies an entity (division / unit / dept).
    """
    parts = [row.get(c, "") for c in (2, 3, 4)]
    return " / ".join(p for p in parts if p)


def _compare_sheets(ws_orig, ws_gen, sheet_name: str) -> dict:
    """
    Compare two worksheets and return a per-sheet comparison dict.
    """
    headers_orig, orig_rows = _read_sheet_data(ws_orig)
    headers_gen,  gen_rows  = _read_sheet_data(ws_gen)
    # Use original headers as the reference for field names
    headers = headers_orig or headers_gen

    orig_keys = [_entity_key(r) for r in orig_rows]
    gen_keys = [_entity_key(r) for r in gen_rows]

    matched_pairs: list[tuple[int, int, float]] = []   # (orig_idx, gen_idx, score)
    used_gen: set[int] = set()

    for oi, ok in enumerate(orig_keys):
        if not ok:
            continue
        # Filter gen_keys to only unused entries
        candidates = [(gk, gi) for gi, gk in enumerate(gen_keys) if gi not in used_gen]
        if not candidates:
            break
        cand_strings = [c[0] for c in candidates]
        cand_indices = [c[1] for c in candidates]

        result = rfprocess.extractOne(
            ok, cand_strings,
            scorer=fuzz.token_sort_ratio,
            score_cutoff=70,
        )
        if result is not None:
            best_str, score, best_local_idx = result
            gi = cand_indices[best_local_idx]
            matched_pairs.append((oi, gi, score))
            used_gen.add(gi)

    matched_orig = {oi for oi, _, _ in matched_pairs}
    matched_gen = {gi for _, gi, _ in matched_pairs}

    missing_rows = [i for i in range(len(orig_rows)) if i not in matched_orig and orig_keys[i]]
    extra_rows = [i for i in range(len(gen_rows)) if i not in matched_gen and gen_keys[i]]

    # Field-level comparison for matched rows
    total_fields = 0
    correct_fields = 0
    top_mismatches: list[dict] = []

    for oi, gi, _ in matched_pairs:
        orig_row = orig_rows[oi]
        gen_row = gen_rows[gi]
        entity = orig_keys[oi]

        # Data columns start at 5 (skip structure cols 1-4)
        all_cols = sorted(
            set(c for c in orig_row if c >= 5) | set(c for c in gen_row if c >= 5)
        )
        for col in all_cols:
            orig_val = orig_row.get(col, "")
            gen_val = gen_row.get(col, "")

            # Both empty → match, no penalty
            if not orig_val and not gen_val:
                continue

            total_fields += 1
            ratio = fuzz.ratio(orig_val, gen_val)
            if ratio >= 85:
                correct_fields += 1
            elif len(top_mismatches) < 5:
                field_name = (
                    headers[col - 1]
                    if col - 1 < len(headers) and headers[col - 1]
                    else f"col_{col}"
                )
                top_mismatches.append({
                    "entity": entity,
                    "field": field_name,
                    "original": orig_val[:80],
                    "generated": gen_val[:80],
                })

    field_accuracy = (correct_fields / total_fields * 100) if total_fields > 0 else 100.0

    return {
        "status": "compared",
        "original_rows": len([k for k in orig_keys if k]),
        "generated_rows": len([k for k in gen_keys if k]),
        "matched_rows": len(matched_pairs),
        "missing_rows": len(missing_rows),
        "extra_rows": len(extra_rows),
        "field_accuracy": round(field_accuracy, 1),
        "top_mismatches": top_mismatches,
    }


def compare_bia_files(original_path, generated_path) -> dict:
    """
    Compare two BIA synthèse Excel files.

    Args:
        original_path: Path to the original (reference) BIA synthèse.
        generated_path: Path to the system-generated BIA synthèse.

    Returns:
        A dict with overall_similarity, overall_field_accuracy, and per-sheet results.
    """
    orig_path = Path(original_path)
    gen_path = Path(generated_path)

    wb_orig = openpyxl.load_workbook(orig_path, data_only=True)
    wb_gen = openpyxl.load_workbook(gen_path, data_only=True)

    orig_sheets = {s for s in wb_orig.sheetnames if s not in _SKIP_SHEETS}
    gen_sheets = {s for s in wb_gen.sheetnames if s not in _SKIP_SHEETS}

    all_sheets = orig_sheets | gen_sheets
    sheets_result: dict[str, dict] = {}

    total_orig_rows = 0
    total_matched_rows = 0
    weighted_field_correct = 0
    weighted_field_total = 0

    for sheet in sorted(all_sheets):
        in_orig = sheet in orig_sheets
        in_gen = sheet in gen_sheets

        if in_orig and in_gen:
            ws_orig = wb_orig[sheet]
            ws_gen = wb_gen[sheet]
            info = _compare_sheets(ws_orig, ws_gen, sheet)
        elif in_orig:
            _, orig_rows = _read_sheet_data(wb_orig[sheet])
            info = {
                "status": "only_in_original",
                "original_rows": len([_entity_key(r) for r in orig_rows if _entity_key(r)]),
                "generated_rows": 0,
                "matched_rows": 0,
                "missing_rows": len([_entity_key(r) for r in orig_rows if _entity_key(r)]),
                "extra_rows": 0,
                "field_accuracy": 0.0,
                "top_mismatches": [],
            }
        else:
            _, gen_rows = _read_sheet_data(wb_gen[sheet])
            info = {
                "status": "only_in_generated",
                "original_rows": 0,
                "generated_rows": len([_entity_key(r) for r in gen_rows if _entity_key(r)]),
                "matched_rows": 0,
                "missing_rows": 0,
                "extra_rows": len([_entity_key(r) for r in gen_rows if _entity_key(r)]),
                "field_accuracy": 0.0,
                "top_mismatches": [],
            }

        sheets_result[sheet] = info

        n_orig = info["original_rows"]
        n_matched = info["matched_rows"]
        total_orig_rows += n_orig
        total_matched_rows += n_matched

        # Weight field accuracy by number of matched rows
        fa = info["field_accuracy"]
        weighted_field_correct += fa * n_matched
        weighted_field_total += n_matched

    overall_similarity = (
        (total_matched_rows / total_orig_rows * 100) if total_orig_rows > 0 else 0.0
    )
    overall_field_accuracy = (
        (weighted_field_correct / weighted_field_total) if weighted_field_total > 0 else 0.0
    )

    return {
        "original_file": orig_path.name,
        "generated_file": gen_path.name,
        "overall_similarity": round(overall_similarity, 1),
        "overall_field_accuracy": round(overall_field_accuracy, 1),
        "sheets": sheets_result,
    }


if __name__ == "__main__":
    import json, sys

    if len(sys.argv) != 3:
        print("Usage: python bia_compare.py <original.xlsx> <generated.xlsx>")
        sys.exit(1)

    result = compare_bia_files(sys.argv[1], sys.argv[2])
    print(json.dumps(result, ensure_ascii=False, indent=2))
