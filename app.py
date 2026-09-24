"""
BIA Web App — FastAPI backend

Endpoints:
  POST /api/process          – Fill the Synthèse BIA from individual fiches (.docx)
  POST /api/generate-fiches  – Generate one BIA fiche per department from a
                                recensement file (.xlsx) + BIA template (.docx)
  POST /api/fill-activities  – Inject activities into a generated BIA fiche (.docx)
"""
import asyncio
import io
import queue as _queue_mod
import re
import shutil
import sys
import tempfile
import threading as _threading
import uuid as _uuid_mod
import zipfile
from pathlib import Path
from typing import List, Optional

# Ensure this file's directory is on sys.path so local modules are importable
# regardless of which directory Python was launched from.
sys.path.insert(0, str(Path(__file__).parent))

from fastapi import FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import HTMLResponse, Response, StreamingResponse
from fastapi.staticfiles import StaticFiles

import json as _json
from bia_etl import (extract, load, _fiche_is_empty, bootstrap_synthese,
                     detect_recensement_columns, extract_org_tree,
                     extract_montee_en_charge, extract_rapport_bia,
                     extract_fiche_equipment,
                     read_impact_dmia, write_impact_dmia,
                     read_applications_it, write_applications_it,
                     update_fiche_dmia_table, update_fiche_applications_table,
                     find_unmatched_activities,
                     _strip_accents)
from fiche_generator import generate_all_fiches
from bia_compare import compare_bia_files
from activity_filler import fill_activities as _fill_activities
from fiche_writer import fill_fiche as _fill_fiche
from PIL import Image, ImageOps
import projects_db as _pdb
import risk_analysis as _risk
import risk_components as _risk_comp
import db_backup as _backup
import questionnaires_db as _quest
import orgchart_vision as _ov
import orgchart_ocr as _ocr
import orgchart_tiled as _tiled

app = FastAPI(title="BIA Automatique")

# ── In-memory store for streamed job results ──────────────────────────────────
# Maps job_id → raw .xlsx bytes.  Entries are consumed (popped) on first GET.
_job_store: dict[str, bytes] = {}


def _run_process_sync(
    synthese_bytes: bytes,
    fiche_list: list,          # [(filename, bytes), …]
    rec_bytes: bytes | None,
    llm_model: str | None,
    custom_mapping: dict | None,
    db_project_ids: list,
    event_q: "_queue_mod.Queue",
    job_id: str,
) -> None:
    """
    Full ETL pipeline running synchronously in a background thread.
    Pushes SSE-ready dicts into event_q; caller reads them via an async generator.
    """
    errors: list[str] = []
    unmatched: list[str] = []
    processed = 0
    total = len(fiche_list) + len(db_project_ids)

    def emit(evt: dict) -> None:
        event_q.put(evt)

    with tempfile.TemporaryDirectory() as tmpdir:
        tmp = Path(tmpdir)
        master = tmp / "master.xlsx"
        master.write_bytes(synthese_bytes)
        output = tmp / "output.xlsx"

        # ── Step 1: bootstrap from recensement (optional) ─────────────────────
        if rec_bytes:
            emit({"type": "progress", "current": 0, "total": total,
                  "step": "bootstrap",
                  "message": "Construction de la structure organisationnelle…"})
            rec_path = tmp / "recensement.xlsx"
            rec_path.write_bytes(rec_bytes)
            try:
                bootstrap_synthese(rec_path, master, output, verbose=False,
                                   custom_mapping=custom_mapping)
            except Exception as exc:
                emit({"type": "error", "message": f"Erreur recensement : {exc}"})
                return
        else:
            output.write_bytes(synthese_bytes)

        # ── Step 2: process uploaded fiches ───────────────────────────────────
        for i, (fname, fbytes) in enumerate(fiche_list):
            emit({"type": "progress", "current": i, "total": total,
                  "step": "fiche", "message": f"Traitement : {fname}"})
            fiche_path = tmp / fname
            fiche_path.write_bytes(fbytes)
            try:
                fiche = extract(fiche_path, llm_model=llm_model)
                if _fiche_is_empty(fiche):
                    errors.append(f"{fname} : fiche vide ou illisible")
                    continue
                load(output, fiche, dry_run=False, output_path=output,
                     verbose=False, unmatched=unmatched)
                processed += 1
            except Exception as exc:
                errors.append(f"{fname} : {exc}")

        # ── Step 3: process DB-sourced fiches ─────────────────────────────────
        for i, pid in enumerate(db_project_ids):
            emit({"type": "progress",
                  "current": len(fiche_list) + i, "total": total,
                  "step": "fiche", "message": f"Fiche base de données #{pid}…"})
            try:
                fiche_data = _pdb.get_fiche_docx_form_data(int(pid))
                if fiche_data:
                    meta   = fiche_data.get("_meta", {})
                    folder = _pdb.project_dir(
                        meta.get("sector", "autres"),
                        meta.get("client", "Client"),
                        "fiche",
                    )
                    docx = folder / meta.get("project_name", "")
                    if docx.exists():
                        fiche = extract(docx, llm_model=llm_model)
                        if not _fiche_is_empty(fiche):
                            load(output, fiche, dry_run=False, output_path=output,
                                 verbose=False, unmatched=unmatched)
                            processed += 1
                            continue
                errors.append(f"Projet DB #{pid} : fichier introuvable ou vide")
            except Exception as exc:
                errors.append(f"Projet DB #{pid} : {exc}")

        # Promote unmatched to warnings
        for ent in set(unmatched):
            errors.append(f'Entité absente du recensement : "{ent}"')

        if processed == 0:
            emit({"type": "error",
                  "message": "\n".join(errors) or "Aucune fiche BIA valide fournie."})
            return

        content = output.read_bytes()

    _job_store[job_id] = content
    emit({
        "type":      "done",
        "job_id":    job_id,
        "filename":  "Synthese_BIA_filled.xlsx",
        "processed": processed,
        "total":     total,
        "errors":    errors,
    })

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

BASE_DIR = Path(__file__).parent

# Serve static files using absolute path — works regardless of working directory
app.mount("/static", StaticFiles(directory=str(BASE_DIR / "static")), name="static")


@app.get("/", response_class=HTMLResponse)
async def root():
    return (BASE_DIR / "static" / "index.html").read_text(encoding="utf-8")


@app.get("/projects", response_class=HTMLResponse)
async def projects_page():
    return (BASE_DIR / "static" / "projects.html").read_text(encoding="utf-8")


@app.get("/fiche-selector", response_class=HTMLResponse)
async def fiche_selector():
    return (BASE_DIR / "static" / "fiche-selector.html").read_text(encoding="utf-8")


@app.get("/fiche-editor", response_class=HTMLResponse)
async def fiche_editor_page():
    return (BASE_DIR / "static" / "fiche-editor.html").read_text(encoding="utf-8")


@app.post("/api/detect-recensement")
async def detect_recensement(
    recensement: UploadFile = File(..., description="Fiche de recensement (.xlsx)"),
):
    """
    Inspect a recensement file and return the auto-detected column mapping
    with a confidence rating.  Used by the UI to show/hide the mapping panel.
    """
    with tempfile.TemporaryDirectory() as tmpdir:
        tmp = Path(tmpdir)
        rec_path = tmp / (recensement.filename or "recensement.xlsx")
        rec_path.write_bytes(await recensement.read())
        try:
            result = detect_recensement_columns(rec_path)
            return result
        except Exception as exc:
            raise HTTPException(status_code=422, detail=str(exc))


@app.post("/api/process")
@app.post("/api/fill-synthese")   # alias used by the frontend
async def process_bia(
    synthese: UploadFile = File(..., description="Synthèse BIA master template (.xlsx)"),
    fiches: List[UploadFile] = File(None, description="Fiches BIA (.docx)"),
    recensement: UploadFile = File(None, description="Fiche de recensement (.xlsx) — optional, seeds org structure"),
    llm_model: str = Form("", description="Ollama model name (optional, e.g. 'qwen2.5:3b')"),
    column_mapping: str = Form("", description="JSON column mapping override, e.g. '{\"division\":2,\"unite\":5,\"departement\":8}'"),
    project_ids: str = Form("", description="JSON array of DB project IDs to include as fiches"),
    use_llm: str = Form("", description="Ignored — presence of llm_model controls LLM use"),
):
    """
    Two-step workflow:
      Step 1 (optional): if recensement is provided, seed the synthèse with the
                         client's org structure (Division/Unité/Département rows).
      Step 2: fill the seeded synthèse with data extracted from each BIA fiche.

    Accepts fiches either as uploaded files OR as DB project IDs (or both).
    Without recensement, falls back to the original single-step behaviour
    (synthèse must already have dept rows pre-filled).
    """
    with tempfile.TemporaryDirectory() as tmpdir:
        tmp = Path(tmpdir)

        # Save the master template
        template_bytes = await synthese.read()
        master = tmp / "master.xlsx"
        master.write_bytes(template_bytes)

        # Parse optional column mapping override
        custom_mapping: dict | None = None
        if column_mapping.strip():
            try:
                raw_map = _json.loads(column_mapping)
                # Normalise: convert string keys/values to expected types
                custom_mapping = {
                    k: (int(v) if v is not None else None)
                    for k, v in raw_map.items()
                    if k in ("division", "unite", "departement")
                }
            except Exception:
                pass  # Invalid JSON — fall back to auto-detection

        # Step 1: bootstrap from recensement (if provided)
        output = tmp / "output.xlsx"
        if recensement and recensement.filename:
            rec_path = tmp / recensement.filename
            rec_path.write_bytes(await recensement.read())
            try:
                bootstrap_synthese(rec_path, master, output, verbose=False,
                                   custom_mapping=custom_mapping)
            except Exception as exc:
                raise HTTPException(
                    status_code=422,
                    detail=f"Erreur lors de la lecture du recensement: {exc}",
                )
        else:
            # No recensement — use template as-is (original behaviour)
            output.write_bytes(template_bytes)

        errors: list[str] = []
        unmatched_entities: list[str] = []   # entities not found in synthèse
        processed = 0

        # ── Resolve DB-based fiches (project_ids) ─────────────────────────────
        db_fiche_paths: list[Path] = []
        if project_ids.strip():
            try:
                ids = _json.loads(project_ids)
                for pid in ids:
                    try:
                        fiche_data = _pdb.get_fiche_docx_form_data(int(pid))
                        if fiche_data is None:
                            errors.append(f"Projet DB #{pid} introuvable — ignoré")
                            continue
                        meta = fiche_data.get("_meta", {})
                        sector = meta.get("sector", "autres")
                        client = meta.get("client", "Client")
                        pname  = meta.get("project_name", "")
                        folder = _pdb.project_dir(sector, client, "fiche")
                        docx   = folder / pname
                        if not docx.exists():
                            # Try case-insensitive search
                            matches = list(folder.glob("*.docx"))
                            found   = next((f for f in matches
                                            if f.name.lower() == pname.lower()), None)
                            if found:
                                docx = found
                            else:
                                errors.append(f"Fichier introuvable pour projet DB #{pid} ({pname}) — ignoré")
                                continue
                        db_fiche_paths.append(docx)
                    except Exception as exc:
                        errors.append(f"Erreur projet DB #{pid}: {exc}")
            except Exception:
                pass  # Invalid JSON for project_ids — ignore silently

        # Build the combined list of fiche paths to process
        upload_list = list(fiches) if fiches else []
        if not upload_list and not db_fiche_paths:
            raise HTTPException(
                status_code=422,
                detail="Aucune fiche BIA fournie. Chargez des fichiers .docx ou sélectionnez des projets depuis la base.",
            )

        # ── Process uploaded .docx files ──────────────────────────────────────
        for upload in upload_list:
            if not upload.filename.lower().endswith(".docx"):
                errors.append(f"{upload.filename}: fichier ignoré (pas un .docx)")
                continue

            fiche_path = tmp / upload.filename
            fiche_path.write_bytes(await upload.read())

            try:
                fiche = extract(fiche_path, llm_model=llm_model.strip() or None)

                if _fiche_is_empty(fiche):
                    errors.append(
                        f"{upload.filename}: fiche vide ou illisible — aucune donnée extraite"
                    )
                    continue

                # output_path=output → overwrites in place so the next fiche
                # reads the already-updated file (sequential accumulation)
                load(output, fiche, dry_run=False, output_path=output, verbose=False,
                     unmatched=unmatched_entities)
                processed += 1
            except Exception as exc:
                errors.append(f"{upload.filename}: {exc}")

        # ── Process DB-sourced .docx files ────────────────────────────────────
        for docx_path in db_fiche_paths:
            try:
                fiche = extract(docx_path, llm_model=llm_model.strip() or None)
                if _fiche_is_empty(fiche):
                    errors.append(f"{docx_path.name}: fiche vide ou illisible — aucune donnée extraite")
                    continue
                load(output, fiche, dry_run=False, output_path=output, verbose=False,
                     unmatched=unmatched_entities)
                processed += 1
            except Exception as exc:
                errors.append(f"{docx_path.name}: {exc}")

        # Promote unmatched entities to user-visible warnings
        if unmatched_entities:
            seen = set()
            for ent in unmatched_entities:
                if ent not in seen:
                    errors.append(
                        f"Entite absente du recensement - ignoree : \"{ent}\" "
                        f"(ajouter cette entite au fichier de recensement)"
                    )
                    seen.add(ent)

        if processed == 0:
            raise HTTPException(
                status_code=422,
                detail=errors or ["Aucune fiche .docx valide fournie."],
            )

        content = output.read_bytes()

    return Response(
        content=content,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": 'attachment; filename="Synthese_BIA_filled.xlsx"',
            "X-Processed-Count": str(processed),
            "X-Error-Count": str(len(errors)),
            "X-Errors": _safe_header(" | ".join(errors)) if errors else "",
            "Access-Control-Expose-Headers": "X-Processed-Count, X-Error-Count, X-Errors",
        },
    )


@app.post("/api/process-stream")
async def process_bia_stream(
    synthese:     UploadFile = File(..., description="Synthèse BIA master template (.xlsx)"),
    fiches:       List[UploadFile] = File(None,  description="Fiches BIA (.docx)"),
    recensement:  UploadFile = File(None,  description="Fiche de recensement (.xlsx)"),
    llm_model:    str = Form(""),
    column_mapping: str = Form(""),
    project_ids:  str = Form(""),
    use_llm:      str = Form(""),
):
    """
    Streaming version of /api/process.
    Returns an SSE stream of progress events, then stores the result under a
    job_id that the client retrieves via GET /api/process-result/{job_id}.
    """
    # Read all files upfront (must be done in the async context before handing
    # off to the sync worker thread).
    synthese_bytes = await synthese.read()
    fiche_list = [(f.filename, await f.read()) for f in (fiches or [])]
    rec_bytes  = (await recensement.read()) if (recensement and recensement.filename) else None

    custom_mapping: dict | None = None
    if column_mapping.strip():
        try:
            raw = _json.loads(column_mapping)
            custom_mapping = {k: (int(v) if v is not None else None)
                              for k, v in raw.items()
                              if k in ("division", "unite", "departement")}
        except Exception:
            pass

    db_ids: list = []
    if project_ids.strip():
        try:
            db_ids = [int(x) for x in _json.loads(project_ids)]
        except Exception:
            pass

    llm = llm_model.strip() or None
    job_id   = _uuid_mod.uuid4().hex[:16]
    event_q  = _queue_mod.Queue()

    _threading.Thread(
        target=_run_process_sync,
        args=(synthese_bytes, fiche_list, rec_bytes, llm,
              custom_mapping, db_ids, event_q, job_id),
        daemon=True,
    ).start()

    async def generate():
        while True:
            try:
                event = event_q.get_nowait()
                yield f"data: {_json.dumps(event, ensure_ascii=False)}\n\n"
                if event.get("type") in ("done", "error"):
                    break
            except _queue_mod.Empty:
                await asyncio.sleep(0.05)
                yield ": \n\n"   # SSE keep-alive comment

    return StreamingResponse(
        generate(),
        media_type="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


@app.get("/api/process-result/{job_id}")
async def get_process_result(job_id: str):
    """Download the .xlsx result produced by /api/process-stream."""
    data = _job_store.pop(job_id, None)
    if not data:
        raise HTTPException(status_code=404,
                            detail="Résultat introuvable ou déjà téléchargé.")
    return Response(
        content=data,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="Synthese_BIA_filled.xlsx"'},
    )


@app.get("/api/ollama-status")
async def ollama_status(model: str = "qwen2.5:3b"):
    """Check if Ollama is running and model is available."""
    try:
        from llm_fallback import is_ollama_available
        available = is_ollama_available(model)
        return {"available": available, "model": model}
    except Exception as exc:
        return {"available": False, "model": model, "error": str(exc)}


@app.post("/api/generate-fiches")
async def generate_fiches(
    recensement: UploadFile = File(..., description="Fiche de recensement (.xlsx)"),
    template: UploadFile = File(..., description="BIA fiche template (.docx)"),
    version: str = Form("2.0", description="Document version: '1.0' or '2.0'"),
    openai_key: str = Form("", description="OpenAI API key (optional, overrides env var)"),
    client_name: str = Form("", description="Client name (optional; skips AI name detection when provided)"),
):
    """
    Accepts a recensement Excel file + a BIA fiche template (.docx).
    Uses GPT-4o Vision to detect the client name and logo from the xlsx.
    Generates one filled BIA fiche per department / structure.
    Returns a ZIP archive containing all generated .docx files.
    """
    if not recensement.filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=422, detail="Le fichier de recensement doit être un .xlsx")
    if not template.filename.lower().endswith(".docx"):
        raise HTTPException(status_code=422, detail="Le modèle BIA doit être un .docx")
    if version not in ("1.0", "2.0"):
        version = "2.0"

    with tempfile.TemporaryDirectory() as tmpdir:
        tmp = Path(tmpdir)

        xlsx_path = tmp / recensement.filename
        xlsx_path.write_bytes(await recensement.read())

        tmpl_path = tmp / template.filename
        tmpl_path.write_bytes(await template.read())

        out_dir = tmp / "fiches"

        generated, errors = generate_all_fiches(
            xlsx_path, tmpl_path, out_dir,
            version=version,
            openai_api_key=openai_key or None,
            client_name=client_name.strip() or None,
        )

        if not generated:
            raise HTTPException(
                status_code=422,
                detail=errors or ["Aucune structure trouvée dans le fichier de recensement."],
            )

        # Pack all generated fiches into a single ZIP
        zip_buf = io.BytesIO()
        with zipfile.ZipFile(zip_buf, "w", zipfile.ZIP_DEFLATED) as zf:
            for fiche_path in generated:
                zf.write(fiche_path, fiche_path.name)
        zip_bytes = zip_buf.getvalue()

    # Extract client name from the first generated filename for the response
    client_name = "Client"
    if generated:
        client_name = generated[0].name.split(" - MCO")[0]

    safe_zip_name = f"Fiches_BIA_{_safe_zip_name(client_name)}_V{version}.zip"

    return Response(
        content=zip_bytes,
        media_type="application/zip",
        headers={
            "Content-Disposition": f'attachment; filename="{safe_zip_name}"',
            "X-Generated-Count": str(len(generated)),
            "X-Error-Count": str(len(errors)),
            "X-Errors": _safe_header(" | ".join(errors)) if errors else "",
            "X-Client-Name": _safe_header(client_name),
            "Access-Control-Expose-Headers": (
                "X-Generated-Count, X-Error-Count, X-Errors, X-Client-Name"
            ),
        },
    )


@app.get("/api/orgchart/models-vision")
async def orgchart_vision_models():
    """Return locally available Ollama vision models."""
    return {"models": _ov.list_vision_models()}


@app.post("/api/orgchart/analyze")
async def orgchart_analyze(
    image: UploadFile = File(..., description="Org chart image (PNG/JPG)"),
    model: str = Form(..., description="Ollama vision model name"),
):
    """Analyze an org chart image with a local vision model. Returns detected entities with hierarchy."""
    allowed = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp"}
    suffix = Path(image.filename).suffix.lower()
    if suffix not in allowed:
        raise HTTPException(
            status_code=422,
            detail=f"Format non supporté '{suffix}'. Utilisez PNG ou JPG."
        )
    image_bytes = await image.read()
    if len(image_bytes) > 15 * 1024 * 1024:
        raise HTTPException(status_code=422, detail="Image trop grande (max 15 Mo).")
    try:
        entities = _ov.analyze_orgchart(image_bytes, model)
    except RuntimeError as e:
        raise HTTPException(status_code=502, detail=str(e))
    return {"entities": entities, "count": len(entities)}


@app.get("/api/orgchart/ocr-available")
async def orgchart_ocr_available():
    """Check whether EasyOCR is installed."""
    return {"available": _ocr.ocr_available()}


@app.post("/api/orgchart/analyze-ocr")
async def orgchart_analyze_ocr(
    image: UploadFile = File(..., description="Org chart image (PNG/JPG)"),
):
    """Analyze an org chart image with OCR + layout analysis (no Ollama needed)."""
    allowed = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp"}
    suffix = Path(image.filename).suffix.lower()
    if suffix not in allowed:
        raise HTTPException(status_code=422, detail=f"Format non supporté '{suffix}'.")
    image_bytes = await image.read()
    if len(image_bytes) > 15 * 1024 * 1024:
        raise HTTPException(status_code=422, detail="Image trop grande (max 15 Mo).")
    if not _ocr.ocr_available():
        raise HTTPException(status_code=503,
            detail="EasyOCR non installé. Exécutez : pip install easyocr")
    try:
        entities = _ocr.analyze_orgchart_ocr(image_bytes)
    except RuntimeError as e:
        raise HTTPException(status_code=502, detail=str(e))
    return {"entities": entities, "count": len(entities)}


@app.post("/api/orgchart/analyze-hybrid")
async def orgchart_analyze_hybrid(
    image: UploadFile = File(..., description="Org chart image (PNG/JPG)"),
    model: str = Form(..., description="Ollama vision model name"),
):
    """Hybrid: OCR extracts entity names, VLM organises hierarchy. Best of both worlds."""
    allowed = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp"}
    suffix = Path(image.filename).suffix.lower()
    if suffix not in allowed:
        raise HTTPException(status_code=422, detail=f"Format non supporté '{suffix}'.")
    image_bytes = await image.read()
    if len(image_bytes) > 15 * 1024 * 1024:
        raise HTTPException(status_code=422, detail="Image trop grande (max 15 Mo).")
    if not _ocr.ocr_available():
        raise HTTPException(status_code=503,
            detail="EasyOCR non installé. Exécutez : pip install easyocr")
    try:
        result = _ocr.analyze_orgchart_hybrid(image_bytes, model)
    except RuntimeError as e:
        raise HTTPException(status_code=502, detail=str(e))
    entities = result.get("entities", [])
    levels   = result.get("levels",   [])
    return {"entities": entities, "levels": levels, "count": len(entities)}


@app.get("/api/orgchart/glmocr-available")
async def orgchart_glmocr_available():
    """Check whether glm-ocr:latest is loaded in Ollama."""
    return {"available": _ocr.glmocr_available()}


@app.post("/api/orgchart/analyze-glmocr")
async def orgchart_analyze_glmocr(
    image: UploadFile = File(..., description="Diagram image (PNG/JPG/WEBP)"),
    model: str = Form("", description="Optional vision model for hierarchy tree (Ollama)"),
):
    """Analyse a hierarchical diagram with GLM-OCR + optional VLM hierarchy tree."""
    allowed = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp"}
    suffix  = Path(image.filename).suffix.lower()
    if suffix not in allowed:
        raise HTTPException(status_code=422, detail=f"Format non supporté '{suffix}'.")
    image_bytes = await image.read()
    if len(image_bytes) > 15 * 1024 * 1024:
        raise HTTPException(status_code=422, detail="Image trop grande (max 15 Mo).")
    try:
        result = _ocr.analyze_orgchart_glmocr(
            image_bytes,
            vlm_model=model.strip() or None,
        )
    except RuntimeError as e:
        raise HTTPException(status_code=502, detail=str(e))
    entities = result.get("entities", [])
    levels   = result.get("levels",   [])
    tree     = result.get("tree",     None)
    return {"entities": entities, "levels": levels, "tree": tree, "count": len(entities)}


@app.post("/api/orgchart/analyze-tiled")
async def orgchart_analyze_tiled(
    image: UploadFile = File(..., description="Org chart image (PNG/JPG)"),
    model: str = Form("qwen2.5vl:3b", description="Ollama vision model name"),
):
    """Analyse avancée : découpage en bandes + extraction IA par bande + liaison hiérarchique bande par bande."""
    allowed = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp"}
    suffix = Path(image.filename).suffix.lower()
    if suffix not in allowed:
        raise HTTPException(status_code=422, detail=f"Format non supporté '{suffix}'.")
    image_bytes = await image.read()
    if len(image_bytes) > 15 * 1024 * 1024:
        raise HTTPException(status_code=422, detail="Image trop grande (max 15 Mo).")
    try:
        result = _tiled.analyze_orgchart_tiled(image_bytes, model)
    except RuntimeError as e:
        raise HTTPException(status_code=502, detail=str(e))
    entities = result.get("entities", [])
    levels   = result.get("levels",   [])
    return {"entities": entities, "levels": levels, "count": len(entities)}


@app.post("/api/orgchart/analyze-tiled-stream")
async def orgchart_analyze_tiled_stream(
    image: UploadFile = File(..., description="Org chart image (PNG/JPG)"),
    model: str = Form("qwen2.5vl:3b", description="Ollama vision model name"),
):
    """
    Streaming version of /api/orgchart/analyze-tiled — SSE progress events
    (one per band extracted, one per band-pair linked) followed by a final
    "done" event carrying the entities/levels, or an "error" event.
    """
    allowed = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp"}
    suffix = Path(image.filename).suffix.lower()
    if suffix not in allowed:
        raise HTTPException(status_code=422, detail=f"Format non supporté '{suffix}'.")
    image_bytes = await image.read()
    if len(image_bytes) > 15 * 1024 * 1024:
        raise HTTPException(status_code=422, detail="Image trop grande (max 15 Mo).")

    event_q: "_queue_mod.Queue" = _queue_mod.Queue()

    def worker() -> None:
        def emit_progress(evt: dict) -> None:
            event_q.put({"type": "progress", **evt})

        try:
            result = _tiled.analyze_orgchart_tiled(image_bytes, model, on_progress=emit_progress)
            event_q.put({
                "type": "done",
                "entities": result.get("entities", []),
                "levels": result.get("levels", []),
                "count": len(result.get("entities", [])),
            })
        except RuntimeError as e:
            event_q.put({"type": "error", "message": str(e)})
        except Exception as e:
            event_q.put({"type": "error", "message": f"Erreur inattendue : {e}"})

    _threading.Thread(target=worker, daemon=True).start()

    async def generate():
        while True:
            try:
                event = event_q.get_nowait()
                yield f"data: {_json.dumps(event, ensure_ascii=False)}\n\n"
                if event.get("type") in ("done", "error"):
                    break
            except _queue_mod.Empty:
                await asyncio.sleep(0.05)
                yield ": \n\n"   # SSE keep-alive comment

    return StreamingResponse(
        generate(),
        media_type="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


@app.post("/api/generate-fiches-from-entities")
async def generate_fiches_from_entities(
    template: UploadFile = File(..., description="BIA fiche template (.docx)"),
    entities_json: str = Form(..., description="JSON array of entity dicts"),
    client_name: str = Form("Client", description="Client name for fiche filenames"),
    version: str = Form("2.0", description="Document version"),
):
    """
    Generate one BIA fiche per entity in entities_json.
    entities_json must be: [{"name":"...","niveau1":"...","niveau2":"...","niveau3":"..."}, ...]
    Returns a ZIP of generated .docx files.
    """
    if not template.filename.lower().endswith(".docx"):
        raise HTTPException(status_code=422, detail="Le modèle BIA doit être un .docx")
    if version not in ("1.0", "2.0"):
        version = "2.0"

    try:
        entities = _json.loads(entities_json)
        if not isinstance(entities, list):
            raise ValueError
    except (ValueError, _json.JSONDecodeError):
        raise HTTPException(status_code=422, detail="entities_json invalide — doit être un tableau JSON.")

    structures = [
        {
            "name":      e.get("name", "").strip(),
            "vis_a_vis": "",
            "date":      "",
            "niveau1":   e.get("niveau1", "") or "",
            "niveau2":   e.get("niveau2", "") or "",
            "niveau3":   e.get("niveau3", "") or "",
        }
        for e in entities
        if e.get("name", "").strip()
    ]

    if not structures:
        raise HTTPException(status_code=422, detail="Aucune entité valide fournie.")

    with tempfile.TemporaryDirectory() as tmpdir:
        tmp = Path(tmpdir)
        tmpl_path = tmp / template.filename
        tmpl_path.write_bytes(await template.read())
        out_dir = tmp / "fiches"
        out_dir.mkdir()

        from fiche_generator import generate_fiche, _safe_filename as _sf
        generated: list[Path] = []
        errors: list[str] = []
        safe_client = _sf(client_name.strip() or "Client")

        for s in structures:
            safe_name = _sf(s["name"])
            out_path = out_dir / f"{safe_client} - MCO - BIA - {safe_name} - V{version}.docx"
            try:
                generate_fiche(tmpl_path, s, None, out_path,
                               client_name=client_name.strip() or "Client",
                               version=version)
                generated.append(out_path)
            except Exception as exc:
                errors.append(f"{s['name']}: {exc}")

        if not generated:
            raise HTTPException(status_code=422, detail=errors or ["Aucune fiche générée."])

        zip_buf = io.BytesIO()
        with zipfile.ZipFile(zip_buf, "w", zipfile.ZIP_DEFLATED) as zf:
            for p in generated:
                zf.write(p, p.name)
        zip_bytes = zip_buf.getvalue()

    safe_zip = f"Fiches_BIA_{_safe_zip_name(client_name.strip() or 'Client')}_V{version}.zip"
    return Response(
        content=zip_bytes,
        media_type="application/zip",
        headers={
            "Content-Disposition": f'attachment; filename="{safe_zip}"',
            "X-Generated-Count": str(len(generated)),
            "X-Error-Count":     str(len(errors)),
            "X-Errors":          _safe_header(" | ".join(errors)) if errors else "",
            "Access-Control-Expose-Headers": "X-Generated-Count, X-Error-Count, X-Errors",
        },
    )


@app.post("/api/orgchart")
async def orgchart(
    synthese: UploadFile = File(..., description="Synthèse BIA (.xlsx)"),
):
    """
    Parse a synthèse BIA and return the organisation hierarchy as a
    nested tree JSON ready for D3.hierarchy().
    """
    with tempfile.TemporaryDirectory() as tmpdir:
        tmp = Path(tmpdir)
        path = tmp / (synthese.filename or "synthese.xlsx")
        path.write_bytes(await synthese.read())
        try:
            tree = extract_org_tree(path)
            return tree
        except Exception as exc:
            raise HTTPException(status_code=422, detail=str(exc))


@app.post("/api/rapport-bia")
async def rapport_bia_endpoint(
    synthese:  UploadFile = File(..., description="Synthèse BIA principale (.xlsx)"),
    synthese2: UploadFile = File(None, description="Synthèse BIA de comparaison (.xlsx) — optionnel"),
):
    """Return all data needed for the rapport BIA charts (recovery curve, HR, equipment)."""
    with tempfile.TemporaryDirectory() as tmpdir:
        tmp  = Path(tmpdir)
        p1   = tmp / (synthese.filename or "synthese.xlsx")
        p1.write_bytes(await synthese.read())
        try:
            data1 = extract_rapport_bia(p1)
        except Exception as exc:
            raise HTTPException(status_code=422, detail=str(exc))

        data2 = None
        if synthese2 and synthese2.filename:
            p2 = tmp / ("cmp_" + (synthese2.filename or "synthese2.xlsx"))
            p2.write_bytes(await synthese2.read())
            try:
                data2 = extract_rapport_bia(p2)
            except Exception:
                data2 = None

        return {"main": data1, "compare": data2}


@app.post("/api/rapport-bia/slides-data")
async def rapport_bia_slides_data(
    synthese: UploadFile = File(..., description="Synthèse BIA (.xlsx)"),
):
    """Return all structured data needed to render the BIA slide gallery."""
    from bia_slides_data import extract_slides_data
    xlsx_bytes = await synthese.read()
    try:
        return extract_slides_data(xlsx_bytes)
    except Exception as exc:
        raise HTTPException(status_code=422, detail=str(exc))


@app.post("/api/rapport-bia/generate-pptx")
async def rapport_bia_generate_pptx(
    synthese: UploadFile = File(..., description="Synthèse BIA (.xlsx)"),
):
    """Generate a full BIA Rapport PPTX from a Synthèse BIA xlsx and return it as a download."""
    from bia_report_pptx import generate_pptx
    from fastapi.responses import Response
    xlsx_bytes = await synthese.read()
    try:
        pptx_bytes = generate_pptx(xlsx_bytes)
    except Exception as exc:
        raise HTTPException(status_code=422, detail=str(exc))
    filename = (synthese.filename or "synthese").replace(".xlsx", "") + "_Rapport_BIA.pptx"
    return Response(
        content=pptx_bytes,
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.post("/api/rapport-bia/slides-to-pptx")
async def slides_to_pptx(request: Request):
    """Receive a list of base64 PNG slide images and return a 16:9 PPTX."""
    import base64, io
    from pptx import Presentation
    from pptx.util import Emu
    from fastapi.responses import Response

    body = await request.json()
    images_b64: list[str] = body.get("images", [])
    filename: str = body.get("filename", "Rapport_BIA") + ".pptx"

    # 16:9 slide size: 10 in × 5.625 in
    SLIDE_W = Emu(9144000)   # 10 inches
    SLIDE_H = Emu(5143500)   # 5.625 inches

    prs = Presentation()
    prs.slide_width  = SLIDE_W
    prs.slide_height = SLIDE_H

    blank_layout = prs.slide_layouts[6]  # completely blank

    if not images_b64:
        raise HTTPException(400, "Aucune image reçue — regénérez les slides.")

    # Each image is handled individually so one malformed capture reports which
    # slide it was instead of collapsing the whole request into a bare 500.
    for idx, b64 in enumerate(images_b64, start=1):
        try:
            # Strip data-URI prefix if present
            if "," in b64:
                b64 = b64.split(",", 1)[1]
            img_bytes = base64.b64decode(b64)
            if not img_bytes:
                raise ValueError("image vide")
            slide = prs.slides.add_slide(blank_layout)
            slide.shapes.add_picture(
                io.BytesIO(img_bytes),
                left=0, top=0,
                width=SLIDE_W, height=SLIDE_H,
            )
        except MemoryError:
            raise HTTPException(
                507,
                f"Mémoire insuffisante à la slide {idx}/{len(images_b64)}. "
                "Redémarrez le serveur puis relancez l'export.",
            )
        except Exception as exc:
            raise HTTPException(
                400,
                f"Image illisible à la slide {idx}/{len(images_b64)} : {exc}",
            )

    buf = io.BytesIO()
    try:
        prs.save(buf)
    except MemoryError:
        raise HTTPException(
            507,
            f"Mémoire insuffisante lors de l'assemblage du PPTX ({len(images_b64)} slides). "
            "Redémarrez le serveur puis relancez l'export.",
        )
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.post("/api/extract-fiches-equipment")
async def extract_fiches_equipment_endpoint(
    fiches: list[UploadFile] = File(..., description="Fiches BIA (.docx)"),
    sector: str = Form(""),
    client: str = Form(""),
    store: bool = Form(False),
):
    """
    Extract 'Autres équipements et outils de communication' from one or more fiche .docx files.
    If store=True, also persists each fiche in the DB under sector/client.
    """
    from projects_db import add_project, store_fiche_equipment, init_db
    init_db()
    results = []
    with tempfile.TemporaryDirectory() as tmpdir:
        tmp = Path(tmpdir)
        for fiche in fiches:
            p = tmp / (fiche.filename or "fiche.docx")
            p.write_bytes(await fiche.read())
            try:
                data = extract_fiche_equipment(p)
            except Exception as exc:
                data = {"entity": fiche.filename, "department": "", "items": [], "error": str(exc)}

            if store and data.get("items"):
                try:
                    from bia_etl import extract_dmias_from_fiche
                    activities = extract_dmias_from_fiche(p)
                    pid = add_project(
                        sector=sector or "Autre",
                        client=client or data.get("entity", "Inconnu"),
                        project_name=fiche.filename or "",
                        file_type="fiche",
                        activities=activities,
                    )
                    store_fiche_equipment(pid, data["entity"], data["items"])
                    data["project_id"] = pid
                except Exception:
                    pass

            results.append(data)
    return results


@app.get("/api/fiches-equipment")
async def fiches_equipment_db(sector: str = "", client: str = ""):
    """Return stored fiche equipment from DB, filtered by sector/client."""
    from projects_db import get_equipment_by_filter
    return get_equipment_by_filter(sector=sector, client=client)


@app.get("/api/synthese-projects")
async def synthese_projects_list(sector: str = "", client: str = ""):
    from projects_db import get_synthese_projects
    return get_synthese_projects(sector=sector, client=client)


@app.get("/api/synthese/{project_id}/impact-dmia")
async def get_impact_dmia(project_id: int):
    from projects_db import get_project_file_path
    path = get_project_file_path(project_id)
    if not path:
        raise HTTPException(404, "Fichier synthèse introuvable")
    try:
        return read_impact_dmia(path)
    except Exception as e:
        raise HTTPException(422, str(e))


@app.get("/api/synthese/{project_id}/applications-it")
async def get_applications_it(project_id: int):
    from projects_db import get_project_file_path
    path = get_project_file_path(project_id)
    if not path:
        raise HTTPException(404, "Fichier synthèse introuvable")
    try:
        return read_applications_it(path)
    except Exception as e:
        raise HTTPException(422, str(e))


def _normalize_dmia(value: str) -> str:
    """
    Normalize DMIA value to standard J+N / H+N format.
    Examples: '3J' → 'J+3', '7J' → 'J+7', '4H' → 'H+4', 'J+3' → 'J+3' (unchanged)
    """
    import re
    if not value:
        return value
    v = value.strip()
    # Already correct format
    if re.match(r'^[HhJj]\+\d+', v):
        return v
    # "NJ" → "J+N"
    m = re.match(r'^(\d+)[Jj]$', v)
    if m:
        return f"J+{m.group(1)}"
    # "NH" → "H+N"
    m = re.match(r'^(\d+)[Hh]$', v)
    if m:
        return f"H+{m.group(1)}"
    # "NJ+X" style edge cases — return as-is
    return v


def _refresh_fiche_json_cache(fiche_id: int, fiche_path,
                              dmia_updates: dict = None,
                              app_updates: dict = None) -> None:
    """
    Patch the fiche_editor JSON cache with the exact values that were written
    to the docx, so the web editor immediately reflects the changes.

    dmia_updates: {activity_name_lower: {"dmia_exprimee": str}}
    app_updates:  {app_name_lower: {"criticite", "dmia", "pmdt", "contournement"}}
    """
    import json as _json, re
    from projects_db import _conn, init_db, PROJECTS_DIR
    from bia_etl import _strip_accents
    init_db()

    with _conn() as con:
        row = con.execute(
            "SELECT sector, client, project_name, file_type, form_data_path FROM projects WHERE id=?",
            (fiche_id,)
        ).fetchone()
    if not row:
        return

    # Locate existing JSON or create path
    json_path_str = row["form_data_path"] or ""
    slug = lambda s: re.sub(r'[^\w\-]', '_', s.strip())[:60]

    if json_path_str and Path(json_path_str).exists():
        json_path = Path(json_path_str)
        try:
            form_data = _json.loads(json_path.read_text(encoding="utf-8"))
        except Exception:
            form_data = {}
    else:
        # No existing JSON — try to build one from the docx
        try:
            from projects_db import extract_full_form_from_fiche
            form_data = extract_full_form_from_fiche(Path(fiche_path))
        except Exception:
            form_data = {"activities": [], "applications": []}
        folder = PROJECTS_DIR / slug(row["sector"]) / slug(row["client"]) / "fiche_editor"
        folder.mkdir(parents=True, exist_ok=True)
        json_path = folder / f"form_{fiche_id}.json"

    # ── Patch DMIA values in activities ──────────────────────────────────────
    if dmia_updates:
        def _match_act(name: str) -> "str | None":
            key = _strip_accents(name.lower().strip())
            if key in dmia_updates:
                return key
            # word-overlap fallback
            kw = set(key.split())
            best, best_score = None, 0.0
            for uk in dmia_updates:
                uw = set(uk.split())
                sc = len(kw & uw) / max(len(kw | uw), 1)
                if sc > 0.45 and sc > best_score:
                    best, best_score = uk, sc
            return best

        for act in form_data.get("activities", []):
            m = _match_act(act.get("name", ""))
            if m:
                new_val = dmia_updates[m].get("dmia_exprimee", "")
                if new_val:
                    act["dmia_exprimee"] = _normalize_dmia(new_val)

    # ── Patch application values ──────────────────────────────────────────────
    if app_updates:
        for app in form_data.get("applications", []):
            key = _strip_accents(app.get("application", "").lower().strip())
            if key in app_updates:
                upd = app_updates[key]
                for field in ("criticite", "dmia", "pmdt", "contournement"):
                    if upd.get(field):
                        app[field] = upd[field]

    # ── Save ─────────────────────────────────────────────────────────────────
    form_data.setdefault("_meta", {}).update({
        "project_id": fiche_id, "source": "bds_sync",
        "sector": row["sector"], "client": row["client"],
        "project_name": row["project_name"],
    })
    json_path.write_text(_json.dumps(form_data, ensure_ascii=False, indent=2), encoding="utf-8")

    # Ensure DB has correct path and type
    with _conn() as con:
        con.execute(
            "UPDATE projects SET form_data_path=?, file_type='fiche_editor' WHERE id=?",
            (str(json_path), fiche_id)
        )


@app.post("/api/synthese/{project_id}/save-impact-dmia")
async def save_impact_dmia(project_id: int, request: Request):
    from projects_db import get_project_file_path, get_linked_fiches, PROJECTS_DIR
    data = await request.json()
    edits: list[dict] = data.get("edits", [])

    path = get_project_file_path(project_id)
    if not path:
        raise HTTPException(404, "Fichier synthèse introuvable")

    # Write to Excel
    write_impact_dmia(path, edits)

    # Build update map for fiches — only propagate rows where DMIA Exprimée is set
    dmia_updates = {
        _strip_accents(e.get("activite","").lower().strip()): {
            "dmia_exprimee": _normalize_dmia(e.get("dmia_exprimee","")),
        }
        for e in edits
        if e.get("activite") and e.get("dmia_exprimee","").strip()
    }

    # Propagate to linked fiches
    fiches_updated: list[str] = []
    fiche_errors: list[str] = []
    for fiche in get_linked_fiches(project_id):
        fiche_path = get_project_file_path(fiche["id"])
        if not fiche_path:
            fiche_errors.append(f"{fiche['project_name']}: file not found")
            continue
        try:
            n = update_fiche_dmia_table(fiche_path, dmia_updates)
            if n > 0:
                fiches_updated.append(fiche["project_name"])
                _refresh_fiche_json_cache(fiche["id"], fiche_path, dmia_updates=dmia_updates)
        except Exception as exc:
            fiche_errors.append(f"{fiche['project_name']}: {exc}")

    # Detect unmatched activities for user resolution
    # Trigger when: DMIA changed but no fiches updated, OR any activity edit had no fiche match
    unmatched: list[dict] = []
    changed_edits = [e for e in edits if e.get("activite")]
    if changed_edits and not fiches_updated:
        fiche_paths = []
        for fiche in get_linked_fiches(project_id):
            fp = get_project_file_path(fiche["id"])
            if fp:
                fiche_paths.append({"id": fiche["id"], "project_name": fiche["project_name"], "path": str(fp)})
        # Use DMIA edits if available, otherwise use any changed activity for candidate discovery
        detect_edits = [e for e in edits if e.get("dmia_exprimee","").strip()] or \
                       [{"activite": e["activite"], "dmia_exprimee": ""} for e in changed_edits[:3]]
        try:
            unmatched = find_unmatched_activities(detect_edits, fiche_paths)
        except Exception:
            pass

    return {"status": "ok", "fiches_updated": fiches_updated,
            "fiche_errors": fiche_errors, "unmatched": unmatched}


@app.post("/api/synthese/{project_id}/apply-mapping")
async def apply_mapping(project_id: int, request: Request):
    """
    Apply user-confirmed activity name mappings to linked fiches.
    Body: {"mappings": [{"activite_synth": str, "activite_fiche": str,
                         "fiche_id": int, "dmia_exprimee": str}]}
    """
    from projects_db import get_project_file_path
    data = await request.json()
    mappings: list[dict] = data.get("mappings", [])

    results: list[dict] = []
    for m in mappings:
        fiche_id   = m.get("fiche_id")
        act_fiche  = m.get("activite_fiche", "")
        dmia_val   = m.get("dmia_exprimee", "")
        if not fiche_id or not act_fiche or not dmia_val:
            continue
        fiche_path = get_project_file_path(fiche_id)
        if not fiche_path:
            results.append({"fiche_id": fiche_id, "status": "file not found"})
            continue
        try:
            # Build update dict using the FICHE activity name (not synth name)
            updates = {_strip_accents(act_fiche.lower().strip()): {"dmia_exprimee": _normalize_dmia(dmia_val)}}
            n = update_fiche_dmia_table(fiche_path, updates)
            if n > 0:
                _refresh_fiche_json_cache(fiche_id, fiche_path, dmia_updates=updates)
            results.append({"fiche_id": fiche_id, "activite_fiche": act_fiche,
                            "updated": n > 0})
        except Exception as exc:
            results.append({"fiche_id": fiche_id, "status": str(exc)})

    return {"status": "ok", "results": results}


@app.post("/api/synthese/{project_id}/save-applications-it")
async def save_applications_it(project_id: int, request: Request):
    from projects_db import get_project_file_path, get_linked_fiches
    data = await request.json()
    edits: list[dict] = data.get("edits", [])

    path = get_project_file_path(project_id)
    if not path:
        raise HTTPException(404, "Fichier synthèse introuvable")

    write_applications_it(path, edits)

    app_updates = {
        _strip_accents(e.get("application","").lower().strip()): {
            "criticite":    e.get("criticite",""),
            "dmia":         e.get("dmia",""),
            "pmdt":         e.get("pmdt",""),
            "contournement": e.get("contournement",""),
        }
        for e in edits if e.get("application")
    }

    fiches_updated: list[str] = []
    fiche_errors: list[str] = []
    for fiche in get_linked_fiches(project_id):
        fiche_path = get_project_file_path(fiche["id"])
        if not fiche_path:
            fiche_errors.append(f"{fiche['project_name']}: file not found")
            continue
        try:
            n = update_fiche_applications_table(fiche_path, app_updates)
            if n > 0:
                fiches_updated.append(fiche["project_name"])
                _refresh_fiche_json_cache(fiche["id"], fiche_path, app_updates=app_updates)
        except Exception as exc:
            fiche_errors.append(f"{fiche['project_name']}: {exc}")

    return {"status": "ok", "fiches_updated": fiches_updated, "fiche_errors": fiche_errors}


@app.get("/api/synthese/{project_id}/linked-fiches")
async def get_linked_fiches_endpoint(project_id: int, auto_link: bool = True):
    from projects_db import get_linked_fiches, link_fiche_to_synthese, _conn, init_db
    init_db()

    # Auto-link all fiches of same sector+client if none linked yet
    if auto_link and not get_linked_fiches(project_id):
        with _conn() as con:
            syn = con.execute("SELECT sector, client FROM projects WHERE id=?", (project_id,)).fetchone()
        if syn:
            with _conn() as con:
                fiches = con.execute(
                    "SELECT id FROM projects WHERE sector=? COLLATE NOCASE AND client=? COLLATE NOCASE "
                    "AND file_type IN ('fiche','fiche_editor')",
                    (syn["sector"], syn["client"])
                ).fetchall()
            for f in fiches:
                link_fiche_to_synthese(project_id, f["id"])

    return get_linked_fiches(project_id)


@app.post("/api/synthese/{project_id}/link-fiches")
async def link_fiches_endpoint(project_id: int, request: Request):
    from projects_db import link_fiche_to_synthese, unlink_fiche_from_synthese, get_linked_fiches
    data = await request.json()
    add_ids: list[int] = data.get("add", [])
    remove_ids: list[int] = data.get("remove", [])
    for fid in add_ids:
        link_fiche_to_synthese(project_id, fid)
    for fid in remove_ids:
        unlink_fiche_from_synthese(project_id, fid)
    return get_linked_fiches(project_id)


@app.get("/api/fiche-projects-flat")
async def fiche_projects_flat(sector: str = "", client: str = ""):
    """Return fiche projects as flat list for BDS linking modal, filtered by sector/client."""
    from projects_db import _conn, init_db
    init_db()
    q = "SELECT id, sector, client, project_name FROM projects WHERE file_type IN ('fiche','fiche_editor')"
    params: list = []
    if sector:
        q += " AND sector=? COLLATE NOCASE"; params.append(sector)
    if client:
        q += " AND client=? COLLATE NOCASE"; params.append(client)
    q += " ORDER BY client, project_name"
    with _conn() as con:
        rows = con.execute(q, params).fetchall()
    return [dict(r) for r in rows]


@app.get("/api/download/synthese/{project_id}")
async def download_synthese(project_id: int):
    from projects_db import get_project_file_path
    path = get_project_file_path(project_id)
    if not path or not path.exists():
        raise HTTPException(404, "Fichier introuvable")
    return Response(
        content=path.read_bytes(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{path.name}"'},
    )


@app.get("/api/download/fiche/{project_id}")
async def download_fiche(project_id: int):
    from projects_db import get_project_file_path
    path = get_project_file_path(project_id)
    if not path or not path.exists():
        raise HTTPException(404, "Fichier introuvable")
    return Response(
        content=path.read_bytes(),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{path.name}"'},
    )


@app.get("/api/sectors-clients")
async def sectors_clients():
    """Return available sectors and clients for filter dropdowns."""
    from projects_db import list_sectors_clients
    return list_sectors_clients()


@app.post("/api/index-equipment")
async def index_equipment(sector: str = Form(""), client: str = Form("")):
    """
    Batch-extract equipment from all stored fiche .docx files matching sector/client,
    then persist into fiche_equipment table. Returns count of processed fiches.
    """
    from projects_db import _conn, init_db, project_dir, store_fiche_equipment, PROJECTS_DIR
    init_db()

    # Find all matching project rows with file_type='fiche'
    q = "SELECT id, sector, client, project_name FROM projects WHERE file_type IN ('fiche','fiche_editor')"
    params: list = []
    if sector:
        q += " AND sector = ? COLLATE NOCASE"; params.append(sector)
    if client:
        q += " AND client = ? COLLATE NOCASE"; params.append(client)

    with _conn() as con:
        rows = con.execute(q, params).fetchall()

    processed, skipped = 0, 0
    for row in rows:
        folder = PROJECTS_DIR / row["sector"].strip() / _slug(row["client"]) / "fiche"
        if not folder.exists():
            skipped += 1; continue
        # Try to find the file
        pname = row["project_name"]
        docx_path = None
        exact = folder / pname
        if exact.exists():
            docx_path = exact
        else:
            for f in folder.glob("*.docx"):
                if f.name.lower() == pname.lower():
                    docx_path = f; break
        if docx_path is None:
            skipped += 1; continue
        try:
            data = extract_fiche_equipment(docx_path)
            if data["items"]:
                store_fiche_equipment(row["id"], data["entity"], data["items"])
                processed += 1
        except Exception:
            skipped += 1

    return {"processed": processed, "skipped": skipped}


def _slug(s: str) -> str:
    import re
    return re.sub(r'[^\w\-]', '_', s.strip())[:60]


@app.post("/api/montee-en-charge")
async def montee_en_charge_endpoint(
    synthese: UploadFile = File(..., description="Synthèse BIA (.xlsx)"),
):
    """
    Parse a synthèse BIA and return flat activity/application lists with
    their DMIA values for the Montée en charge diagram.
    """
    with tempfile.TemporaryDirectory() as tmpdir:
        tmp  = Path(tmpdir)
        path = tmp / (synthese.filename or "synthese.xlsx")
        path.write_bytes(await synthese.read())
        try:
            return extract_montee_en_charge(path)
        except Exception as exc:
            raise HTTPException(status_code=422, detail=str(exc))


@app.post("/api/compare")
async def compare_syntheses(
    original: UploadFile = File(..., description="Original BIA synthèse (.xlsx)"),
    generated: UploadFile = File(..., description="System-generated BIA synthèse (.xlsx)"),
):
    """
    Compare two BIA synthèse Excel files and return a structured similarity report.
    """
    with tempfile.TemporaryDirectory() as tmpdir:
        tmp = Path(tmpdir)
        orig_path = tmp / original.filename
        gen_path = tmp / generated.filename
        orig_path.write_bytes(await original.read())
        gen_path.write_bytes(await generated.read())
        try:
            result = compare_bia_files(orig_path, gen_path)
            return result
        except Exception as exc:
            raise HTTPException(status_code=422, detail=str(exc))


@app.post("/api/fill-activities")
async def fill_activities_endpoint(
    fiche: UploadFile = File(..., description="Generated BIA fiche (.docx)"),
    activities: str = Form(..., description='JSON array: [{"name":"…","description":"…"}, …]'),
):
    """
    Inject a list of activities into an already-generated BIA fiche (.docx).

    The following sections are updated automatically:
      • Présentation générale de l'activité  → activity names + descriptions
      • Section 4 – Description et criticité  → one row per activity
      • Section 5.2 – Évaluation des impacts  → one table per activity
      • Section 5.3 – DMIA                    → one row per activity

    Returns the updated .docx file.
    """
    if not fiche.filename.lower().endswith(".docx"):
        raise HTTPException(status_code=422, detail="Le fichier doit être un .docx")

    try:
        activity_list = _json.loads(activities)
        if not isinstance(activity_list, list):
            raise ValueError("activities must be a JSON array")
        for item in activity_list:
            if not isinstance(item, dict) or "name" not in item:
                raise ValueError("Each activity must have a 'name' field")
            item.setdefault("description", "")
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f"Format activities invalide: {exc}")

    with tempfile.TemporaryDirectory() as tmpdir:
        tmp = Path(tmpdir)
        src = tmp / (fiche.filename or "fiche.docx")
        src.write_bytes(await fiche.read())
        out = tmp / f"filled_{fiche.filename or 'fiche.docx'}"
        try:
            _fill_activities(src, activity_list, out)
        except Exception as exc:
            raise HTTPException(status_code=422, detail=str(exc))
        content = out.read_bytes()

    safe_name = _safe_header(out.name)
    return Response(
        content=content,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={
            "Content-Disposition": f'attachment; filename="{safe_name}"',
            "X-Activity-Count": str(len(activity_list)),
            "Access-Control-Expose-Headers": "X-Activity-Count",
        },
    )


@app.post("/api/fill-fiche")
async def fill_fiche_endpoint(
    template: UploadFile = File(..., description="BIA fiche template (.docx)"),
    form_data: str = Form(..., description="JSON with activities + §7.1 data"),
):
    """
    Fill a complete BIA fiche template (.docx) from structured web-editor data.

    Sections auto-filled:
      §3   – Présentation générale de l'activité
      §4   – Description et criticité
      §5.2 – Évaluation des impacts (one table per activity)
      §5.3 – DMIA (Processus | DMIA Exprimée | Premières actions)
      §7.1 – Montée en charge (7 metrics × N activities, DMIA blocking respected)

    Returns the updated .docx file.
    """
    if not template.filename.lower().endswith(".docx"):
        raise HTTPException(status_code=422, detail="Le template doit être un .docx")

    try:
        data = _json.loads(form_data)
        if not isinstance(data, dict):
            raise ValueError("form_data must be a JSON object")
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f"form_data invalide: {exc}")

    with tempfile.TemporaryDirectory() as tmpdir:
        tmp = Path(tmpdir)
        src = tmp / (template.filename or "template.docx")
        src.write_bytes(await template.read())
        out = tmp / f"filled_{template.filename or 'fiche.docx'}"
        try:
            _fill_fiche(src, data, out)
        except Exception as exc:
            raise HTTPException(status_code=422, detail=str(exc))
        content = out.read_bytes()

    n_act = len(data.get("activities", []))
    safe_name = _safe_header(out.name)
    return Response(
        content=content,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={
            "Content-Disposition": f'attachment; filename="{safe_name}"',
            "X-Activity-Count": str(n_act),
            "Access-Control-Expose-Headers": "X-Activity-Count",
        },
    )


@app.post("/api/save-fiches-zip")
async def save_fiches_zip(
    zip_file: UploadFile = File(..., description="ZIP of generated BIA fiches (.zip)"),
    sector:   str = Form("autres",  description="Sector: banques | assurances | autres"),
    client:   str = Form("Client",  description="Client name"),
):
    """
    Extract every .docx file from the ZIP, copy to the project library,
    and register each one in the SQLite database with its activity/DMIA data.
    Returns a summary: {saved, errors, total_files, total_activities}.
    """
    sector = sector.strip() or "autres"
    client = client.strip() or "Client"
    dest_dir = _pdb.project_dir(sector, client, "fiche")

    zip_bytes = await zip_file.read()
    saved:  list[dict] = []
    errors: list[str]  = []

    try:
        with zipfile.ZipFile(io.BytesIO(zip_bytes)) as zf:
            for name in zf.namelist():
                if not name.lower().endswith(".docx"):
                    continue
                fname     = Path(name).name
                docx_data = zf.read(name)

                # Persist to disk
                (dest_dir / fname).write_bytes(docx_data)

                # Extract activities for the DB
                activities: list[dict] = []
                with tempfile.TemporaryDirectory() as tmpdir:
                    tmp_path = Path(tmpdir) / fname
                    tmp_path.write_bytes(docx_data)
                    try:
                        activities = _pdb.extract_dmias_from_fiche(tmp_path)
                    except Exception as exc:
                        errors.append(f"{fname}: extraction error — {exc}")

                pid = _pdb.add_project(
                    sector=sector, client=client,
                    project_name=fname,
                    file_type="fiche",
                    activities=activities,
                )
                saved.append({"file": fname, "project_id": pid,
                              "activities": len(activities)})
    except zipfile.BadZipFile as exc:
        raise HTTPException(status_code=422, detail=f"Fichier ZIP invalide: {exc}")

    return {
        "saved":             saved,
        "errors":            errors,
        "total_files":       len(saved),
        "total_activities":  sum(s["activities"] for s in saved),
    }


# ─────────────────────────────────────────────────────────────────────────────
# PROJECT DATABASE  endpoints
# ─────────────────────────────────────────────────────────────────────────────

@app.get("/api/projects/list")
async def projects_list():
    """Return all stored projects grouped by sector → client."""
    try:
        return _pdb.list_projects()
    except Exception as exc:
        raise HTTPException(status_code=500, detail=str(exc))


@app.post("/api/projects/upload")
async def projects_upload(
    sector:    str = Form(..., description="assurances | bancaire | …"),
    client:    str = Form(..., description="Client name"),
    file_type: str = Form(..., description="fiche | synthese"),
    files: List[UploadFile] = File(..., description="One or more .docx / .xlsx files"),
):
    """
    Save uploaded BIA files to the project library and extract activity/DMIA data.
    """
    if file_type not in ("fiche", "synthese"):
        raise HTTPException(status_code=422, detail="file_type must be 'fiche' or 'synthese'")

    dest_dir = _pdb.project_dir(sector, client, file_type)
    saved: list[dict] = []
    errors: list[str] = []

    with tempfile.TemporaryDirectory() as tmpdir:
        tmp = Path(tmpdir)

        for upload in files:
            fname = upload.filename or "file"
            fpath = tmp / fname
            fpath.write_bytes(await upload.read())

            # Copy to persistent storage
            dest = dest_dir / fname
            import shutil as _shutil
            _shutil.copy2(str(fpath), str(dest))

            # Extract activities/DMIAs
            activities: list[dict] = []
            try:
                if file_type == "fiche" and fname.lower().endswith(".docx"):
                    activities = _pdb.extract_dmias_from_fiche(fpath)
                elif file_type == "synthese" and fname.lower().endswith(".xlsx"):
                    activities = _pdb.extract_dmias_from_synthese(fpath)
            except Exception as exc:
                errors.append(f"{fname}: extraction error — {exc}")

            if activities:
                pid = _pdb.add_project(
                    sector=sector, client=client,
                    project_name=fname,
                    file_type=file_type,
                    activities=activities,
                )
                saved.append({"file": fname, "project_id": pid, "activities": len(activities)})
            else:
                errors.append(f"{fname}: no activities extracted")

    return {
        "saved": saved,
        "errors": errors,
        "total_activities": sum(s["activities"] for s in saved),
    }


@app.delete("/api/projects/{project_id}")
async def projects_delete(project_id: int):
    """Remove a project and all its activity/DMIA records."""
    try:
        _pdb.delete_project(project_id)
        return {"deleted": project_id}
    except Exception as exc:
        raise HTTPException(status_code=500, detail=str(exc))


@app.get("/api/dmia-suggestions")
async def dmia_suggestions(activity: str, sector: str = ""):
    """
    Return DMIA suggestions for an activity name, optionally filtered by sector.
    """
    try:
        return _pdb.get_dmia_suggestions(activity, sector=sector)
    except Exception as exc:
        raise HTTPException(status_code=500, detail=str(exc))


@app.post("/api/dmia-alerts")
async def dmia_alerts(
    payload: str = Form(..., description='JSON: [{activity_name, dmia_exprimee, dmia_minutes}, …]'),
    sector:  str = Form("", description="Optional sector filter"),
):
    """
    Given the current list of activities with their DMIAs, return any
    alerts where a lower DMIA exists in the database for similar activities.
    """
    try:
        activities = _json.loads(payload)
        return _pdb.get_lower_dmia_alerts(activities, sector=sector)
    except Exception as exc:
        raise HTTPException(status_code=422, detail=str(exc))


@app.get("/api/fiche-projects")
async def list_fiche_projects(sector: str = ""):
    """List editor-saved fiches grouped by client, optionally filtered by sector."""
    try:
        return _pdb.list_fiche_editor_projects(sector=sector)
    except Exception as exc:
        raise HTTPException(status_code=500, detail=str(exc))


@app.get("/api/project-form/{project_id}")
async def get_project_form(project_id: int):
    """
    Return the saved form_data JSON for a fiche-editor project.
    For uploaded .docx fiches (file_type='fiche') with no JSON, returns a
    minimal skeleton with the extracted activity/DMIA pairs pre-populated.
    """
    try:
        # Try full editor JSON first
        data = _pdb.load_fiche_form(project_id)
        if data is not None:
            return data

        # Fall back: extract full form data from the .docx for uploaded fiches
        form_data = _pdb.get_fiche_docx_form_data(project_id)
        if form_data is None:
            raise HTTPException(status_code=404, detail="Project not found")
        return form_data

    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=500, detail=str(exc))


@app.post("/api/save-fiche-project")
async def save_fiche_project(
    sector:       str = Form(...),
    client:       str = Form(...),
    project_name: str = Form(...),
    form_data:    str = Form(...),
    project_id:   str = Form(""),
):
    """Create or update a fiche-editor project (form data + activity/DMIA pairs)."""
    try:
        pid  = int(project_id.strip()) if project_id.strip() else None
        data = _json.loads(form_data)
        new_pid = _pdb.save_fiche_form(sector, client, project_name, data, pid)
        return {"project_id": new_pid}
    except Exception as exc:
        raise HTTPException(status_code=422, detail=str(exc))


@app.post("/api/sync-check")
async def sync_check(
    synthese: UploadFile = File(..., description="Synthèse BIA (.xlsx)"),
    project_ids: str = Form("[]", description="JSON array of DB project IDs"),
    fiches: List[UploadFile] = File([], description="Uploaded .docx fiches"),
):
    """
    Compare DMIA values between a synthèse BIA (.xlsx) and a set of fiches BIA (.docx).
    Returns a list of mismatches: activity whose DMIA in the fiche differs from
    'DMIA arbitrée' (or 'DMIA préconisée' if arbitrée is empty) in the synthèse.
    """
    import openpyxl, tempfile, json as _json
    from fiche_writer import parse_dmia_minutes
    from projects_db import extract_full_form_from_fiche, get_fiche_docx_form_data, _conn, init_db

    # ── 1. Read synthèse xlsx ────────────────────────────────────────────────
    synth_bytes = await synthese.read()
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tf:
        tf.write(synth_bytes)
        tf_path = tf.name

    try:
        wb = openpyxl.load_workbook(tf_path, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Impossible de lire la synthèse: {exc}")
    finally:
        import os; os.unlink(tf_path)

    # Scan all sheets for DMIA columns
    # Columns we look for: "Entité"/"Département", "Activité/Processus", "DMIA arbitrée", "DMIA préconisée"
    synthese_rows: list[dict] = []   # {entity, activity, dmia_synth, dmia_min}
    for ws in wb.worksheets:
        headers: dict[str, int] = {}
        for row in ws.iter_rows():
            vals = [str(c.value or "").strip() for c in row]
            if not any(vals):
                continue
            # Detect header row
            if not headers:
                for ci, v in enumerate(vals):
                    vl = v.lower()
                    if any(k in vl for k in ("entité", "entite", "département", "departement", "direction", "service")):
                        headers["entity"] = ci
                    elif any(k in vl for k in ("activité", "activite", "processus")):
                        headers["activity"] = ci
                    elif "arbitr" in vl and "dmia" in vl:
                        headers["dmia_arb"] = ci
                    elif "préconis" in vl and "dmia" in vl:
                        headers["dmia_prec"] = ci
                continue
            if not headers:
                continue
            entity   = vals[headers.get("entity", -1)]   if headers.get("entity")   is not None else ""
            activity = vals[headers.get("activity", -1)] if headers.get("activity") is not None else ""
            dmia_arb  = vals[headers.get("dmia_arb",  -1)] if "dmia_arb"  in headers else ""
            dmia_prec = vals[headers.get("dmia_prec", -1)] if "dmia_prec" in headers else ""
            dmia_val = dmia_arb if dmia_arb else dmia_prec
            if not activity and not entity:
                continue
            dmia_min = parse_dmia_minutes(dmia_val)
            synthese_rows.append({
                "entity":   entity,
                "activity": activity,
                "dmia_synth": dmia_val,
                "dmia_min": dmia_min,
            })

    # ── 2. Collect fiche form data ───────────────────────────────────────────
    fiche_forms: list[dict] = []   # each = {"source": name, "project_id": id|None, "form": form_data}

    # 2a. From uploaded .docx files
    for up in fiches:
        data = await up.read()
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tf:
            tf.write(data)
            tf_path = tf.name
        try:
            form = extract_full_form_from_fiche(Path(tf_path))
            fiche_forms.append({"source": up.filename, "project_id": None, "form": form})
        except Exception:
            pass
        finally:
            import os; os.unlink(tf_path)

    # 2b. From DB project IDs
    try:
        ids = _json.loads(project_ids) if project_ids else []
    except Exception:
        ids = []
    init_db()
    for pid in ids:
        form = get_fiche_docx_form_data(int(pid))
        if form:
            meta = form.get("_meta", {})
            fiche_forms.append({
                "source": meta.get("project_name", f"Project {pid}"),
                "project_id": int(pid),
                "form": form,
            })

    # ── 3. Match and compare ─────────────────────────────────────────────────
    mismatches: list[dict] = []

    def names_match(a: str, b: str) -> bool:
        a, b = a.lower().strip(), b.lower().strip()
        if not a or not b:
            return False
        return a == b or a in b or b in a

    for srow in synthese_rows:
        if srow["dmia_min"] < 0:
            continue  # synthèse has no DMIA to compare

        # Find matching fiche
        for ff in fiche_forms:
            form = ff["form"]
            # Check if entity matches (using suivi.entite or client meta)
            meta = form.get("_meta", {})
            fiche_entity = form.get("suivi", {}).get("entite", "") or meta.get("project_name", "")

            entity_matches = (not srow["entity"]) or names_match(srow["entity"], fiche_entity)
            if not entity_matches:
                continue

            # Find matching activity
            for act in form.get("activities", []):
                if not names_match(srow["activity"], act["name"]) and srow["activity"]:
                    continue
                fiche_dmia = act.get("dmia_exprimee", "")
                fiche_min  = parse_dmia_minutes(fiche_dmia)
                if fiche_min < 0:
                    continue  # fiche has no DMIA for this activity
                # Compare
                if fiche_min != srow["dmia_min"]:
                    mismatches.append({
                        "fiche_source":   ff["source"],
                        "project_id":     ff["project_id"],
                        "entity":         srow["entity"] or fiche_entity,
                        "activity":       srow["activity"] or act["name"],
                        "dmia_synthese":  srow["dmia_synth"],
                        "dmia_fiche":     fiche_dmia,
                        "resolved":       False,
                    })
                break  # matched activity, move to next synthèse row

    return {"mismatches": mismatches, "total": len(mismatches)}


def _safe_zip_name(name: str) -> str:
    return re.sub(r'[^\w\-]', '_', name)[:40]


def _safe_header(value: str) -> str:
    """Encode a string for use in an HTTP header (latin-1 safe)."""
    return value.encode("latin-1", errors="replace").decode("latin-1")


# ══════════════════════════════════════════════════════════════════════════════
# Risk Analysis routes
# ══════════════════════════════════════════════════════════════════════════════

# In-memory store for original Excel bytes (keyed by file_id) so we can
# re-export with updated Plan d'actions without re-uploading each time.
_risk_excel_store: dict[int, bytes] = {}


@app.post("/api/risk/upload")
async def risk_upload(
    file: UploadFile = File(...),
    name: str = Form(""),
    sector: str = Form(""),
    client: str = Form(""),
    sheet: str = Form(""),
):
    raw = await file.read()
    sheet_arg = sheet.strip() or None
    try:
        rows, detected_sheet, sheets, scale = _risk.parse_risk_excel(raw, sheet_arg)
    except Exception as e:
        raise HTTPException(400, f"Impossible de lire le fichier : {e}")

    fname = name.strip() or file.filename or "Analyse des risques"
    file_id = _risk.save_risk_file(fname, sector, client, rows, detected_sheet, scale)
    _risk_excel_store[file_id] = raw

    return {
        "file_id": file_id,
        "sheet": detected_sheet,
        "sheets": sheets,
        "rows": rows,
        "total": len(rows),
        "scale": scale,
        "diverged": sum(1 for r in rows if r.get("diverged")),
    }


@app.get("/api/risk/list")
async def risk_list():
    return _risk.list_risk_files()


@app.get("/api/risk/load/{file_id}")
async def risk_load(file_id: int):
    data = _risk.load_risk_file(file_id)
    if not data:
        raise HTTPException(404, "Fichier introuvable")
    return data


@app.post("/api/risk/save/{file_id}")
async def risk_save(file_id: int, request: Request):
    """
    Persist edited rows, re-scoring them server-side first.

    The browser scores as you type for immediate feedback, but the value that
    reaches the database and the export goes through the same engine that
    parsed the file, so the two can never drift.
    """
    body = await request.json()
    rows = body.get("rows", [])

    existing = _risk.load_risk_file(file_id)
    scale_dict = (existing or {}).get("scale") or {}
    try:
        scale = _risk._scale_from_dict(scale_dict) if scale_dict else None
    except Exception:
        scale = None
    for r in rows:
        _risk.apply_scoring(r, scale)

    _risk.update_risk_file(file_id, rows)
    return {
        "ok": True,
        "rows": rows,
        "diverged": sum(1 for r in rows if r.get("diverged")),
    }


@app.delete("/api/risk/delete/{file_id}")
async def risk_delete(file_id: int):
    _risk.delete_risk_file(file_id)
    _risk_excel_store.pop(file_id, None)
    return {"ok": True}


@app.get("/api/risk/export/{file_id}")
async def risk_export(file_id: int):
    data = _risk.load_risk_file(file_id)
    if not data:
        raise HTTPException(404, "Fichier introuvable")
    original = _risk_excel_store.get(file_id)
    if not original:
        raise HTTPException(400, "Fichier source non disponible — veuillez le re-télécharger.")
    out = _risk.export_risk_excel(original, data["rows"], data["sheet"],
                                  data.get("scale"))
    safe_name = _safe_header(f"{data['name']}_plan_actions.xlsx")
    return Response(
        content=out,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{safe_name}"'},
    )


@app.post("/api/risk/suggest")
async def risk_suggest(request: Request):
    """Stream an AI suggestion for a single risk row via SSE."""
    body = await request.json()
    row   = body.get("row", {})
    model = body.get("model", "")
    if not row or not model:
        raise HTTPException(400, "row et model sont requis")

    def _gen():
        yield from _risk.stream_suggestion(row, model)

    return StreamingResponse(_gen(), media_type="text/event-stream",
                             headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"})


@app.get("/api/risk/models")
async def risk_models():
    return {"models": _risk.list_ollama_models()}


# ══════════════════════════════════════════════════════════════════════════════
# Questionnaires routes (Gestion des risques)
# ══════════════════════════════════════════════════════════════════════════════

@app.get("/api/questionnaires/types")
async def questionnaires_types():
    return _quest.QUESTIONNAIRE_TYPES


@app.get("/api/questionnaires/counts")
async def questionnaires_counts(sector: str = "", client: str = ""):
    return _quest.count_by_type(sector, client)


@app.get("/api/questionnaires/list")
async def questionnaires_list(sector: str = "", client: str = "", qtype: str = ""):
    return _quest.list_questionnaires(sector, client, qtype)


@app.get("/api/questionnaires/template/{qtype}")
async def questionnaires_template(qtype: str):
    data = _quest.get_template(qtype)
    if not data:
        raise HTTPException(404, "Aucun modèle enregistré pour ce type")
    return data


@app.get("/api/questionnaires/load/{qid}")
async def questionnaires_load(qid: int):
    data = _quest.load_questionnaire(qid)
    if not data:
        raise HTTPException(404, "Questionnaire introuvable")
    return data


@app.post("/api/questionnaires/parse")
async def questionnaires_parse(file: UploadFile = File(...), sheet: str = Form("")):
    """Parse-only (no DB write) — used by the Réponses comparison flow to
    preview a client's file before the user decides to save it."""
    raw = await file.read()
    try:
        columns, rows, detected_sheet, sheets = _quest.parse_questionnaire_excel(raw, sheet.strip() or None)
    except Exception as e:
        raise HTTPException(400, f"Impossible de lire le fichier : {e}")
    return {"columns": columns, "rows": rows, "sheet": detected_sheet, "sheets": sheets}


@app.post("/api/questionnaires/upload")
async def questionnaires_upload(
    file: UploadFile = File(...),
    name: str = Form(""),
    sector: str = Form(""),
    client: str = Form(""),
    qtype: str = Form(...),
    sheet: str = Form(""),
    as_template: str = Form(""),
):
    raw = await file.read()
    sheet_arg = sheet.strip() or None
    try:
        columns, rows, detected_sheet, sheets = _quest.parse_questionnaire_excel(raw, sheet_arg)
    except Exception as e:
        raise HTTPException(400, f"Impossible de lire le fichier : {e}")

    fname = name.strip() or file.filename or "Questionnaire"
    is_template = as_template.strip().lower() in ("1", "true", "yes", "on")
    try:
        qid = _quest.save_questionnaire(fname, sector, client, qtype, columns, rows, detected_sheet, raw, is_template)
    except ValueError as e:
        raise HTTPException(400, str(e))

    return {"id": qid, "sheet": detected_sheet, "sheets": sheets, "columns": columns, "rows": rows, "total": len(rows)}


@app.post("/api/questionnaires/create-from-template")
async def questionnaires_create_from_template(request: Request):
    body = await request.json()
    qtype   = body.get("qtype", "")
    sector  = body.get("sector", "")
    client  = body.get("client", "")
    name    = body.get("name") or "Nouveau questionnaire"
    columns = body.get("columns", [])
    rows    = body.get("rows", [])
    try:
        qid = _quest.create_from_template(qtype, sector, client, name, columns, rows)
    except ValueError as e:
        raise HTTPException(400, str(e))
    return {"id": qid}


@app.post("/api/questionnaires/save-template")
async def questionnaires_save_template_direct(request: Request):
    """Save (or update) the reference template for a type — no prior upload required."""
    body = await request.json()
    qtype   = body.get("qtype", "")
    name    = body.get("name") or "Questionnaire de référence"
    columns = body.get("columns", [])
    rows    = body.get("rows", [])
    try:
        qid = _quest.save_template_direct(qtype, columns, rows, name)
    except ValueError as e:
        raise HTTPException(400, str(e))
    return {"id": qid}


@app.post("/api/questionnaires/save-direct")
async def questionnaires_save_direct(request: Request):
    """Save a client questionnaire without requiring a pre-existing template."""
    body = await request.json()
    qtype   = body.get("qtype", "")
    sector  = body.get("sector", "")
    client  = body.get("client", "")
    name    = body.get("name") or "Questionnaire client"
    columns = body.get("columns", [])
    rows    = body.get("rows", [])
    try:
        qid = _quest.save_questionnaire_direct(qtype, sector, client, name, columns, rows)
    except ValueError as e:
        raise HTTPException(400, str(e))
    return {"id": qid}


@app.post("/api/questionnaires/save/{qid}")
async def questionnaires_save(qid: int, request: Request):
    body = await request.json()
    columns = body.get("columns", [])
    rows = body.get("rows", [])
    name = body.get("name")
    _quest.update_questionnaire(qid, columns, rows, name)
    return {"ok": True}


@app.post("/api/questionnaires/save-as-template/{qid}")
async def questionnaires_save_as_template(qid: int):
    try:
        new_id = _quest.set_as_template(qid)
    except ValueError as e:
        raise HTTPException(404, str(e))
    return {"ok": True, "template_id": new_id}


@app.delete("/api/questionnaires/delete/{qid}")
async def questionnaires_delete(qid: int):
    _quest.delete_questionnaire(qid)
    return {"ok": True}


@app.get("/api/questionnaires/export/{qid}")
async def questionnaires_export(qid: int):
    data = _quest.load_questionnaire(qid)
    if not data:
        raise HTTPException(404, "Questionnaire introuvable")
    original = _quest.get_original_bytes(qid)
    if not original:
        raise HTTPException(400, "Fichier source introuvable sur le disque.")
    out = _quest.export_questionnaire_excel(original, data["columns"], data["rows"], data["sheet_name"])
    safe_name = _safe_header(f"{data['name']}.xlsx")
    return Response(
        content=out,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{safe_name}"'},
    )


@app.get("/api/backup/status")
async def backup_status():
    """Liste des sauvegardes disponibles et état de la base courante."""
    return _backup.status()


@app.post("/api/backup/now")
async def backup_now():
    """Force une sauvegarde immédiate, sans attendre la détection automatique."""
    return _backup.sync(reason="manuel (interface)")


@app.post("/api/backup/restore/{backup_name}")
async def backup_restore(backup_name: str):
    """Restaure une sauvegarde ; la base courante est mise de côté avant."""
    if "/" in backup_name or "\\" in backup_name or ".." in backup_name:
        raise HTTPException(400, "Nom de sauvegarde invalide")
    res = _backup.restore(backup_name)
    if res.get("status") == "erreur":
        raise HTTPException(404, res["raison"])
    return res


@app.on_event("startup")
async def _startup_backup():
    """
    Sauvegarde au démarrage, puis surveillance continue : toute écriture dans
    projects.db est répercutée sur les copies dans les secondes qui suivent.
    """
    try:
        res = _backup.sync(reason="démarrage du serveur")
        if res.get("status") == "ok":
            print(f"[sauvegarde] {res['fichier']} — référence {res['reference']}")
        _backup.start_watcher()
        print(f"[sauvegarde] surveillance active — {_backup.BACKUP_DIR}")
    except Exception as exc:
        print(f"[sauvegarde] échec au démarrage : {exc}")


if __name__ == "__main__":
    import uvicorn
    # reload=False avoids the Windows multiprocessing crash where the
    # reloader worker can't find app.py when cwd differs from the file's location
    uvicorn.run("app:app", host="0.0.0.0", port=8000, reload=False)

# ══════════════════════════════════════════════════════════════════════════════
#  Analyse des risques — relevé photo par étage
#
#  Hiérarchie : évaluation → étage → composant → note → photos
#  Chaque photo est étiquetée (étage, composant, état, niveau de risque) afin
#  qu'un agent puisse la rattacher au bon constat lors de la synthèse.
# ══════════════════════════════════════════════════════════════════════════════

_RA_PHOTO_MAX_BYTES = 15 * 1024 * 1024      # 15 Mo par photo en entrée
_RA_MAX_EDGE        = 2048                  # côté max après compression
_RA_THUMB_EDGE      = 320
_RA_CONDITIONS      = {"bon", "moyen", "mauvais", "non_renseigne"}
_RA_RISK_LEVELS     = {"faible", "moyen", "eleve", "critique", "non_renseigne"}


def _ra_photo_dir(assessment: dict) -> Path:
    d = _pdb.project_dir(assessment["sector"], assessment["client"],
                         "analyse_risques") / str(assessment["id"]) / "photos"
    (d / "vignettes").mkdir(parents=True, exist_ok=True)
    return d


def _ra_exif(img) -> dict:
    """Horodatage et orientation EXIF, quand l'appareil les fournit."""
    meta = {"timestamp": None, "orientation": 1}
    try:
        exif = img.getexif()
        if exif:
            # 36867 = DateTimeOriginal, 306 = DateTime, 274 = Orientation
            raw = exif.get(36867) or exif.get(306)
            if raw:
                meta["timestamp"] = str(raw).replace(":", "-", 2)
            meta["orientation"] = int(exif.get(274) or 1)
    except Exception:
        pass
    return meta


@app.get("/api/risk-assessment/building-types")
async def ra_building_types():
    """Types de bâtiment : catalogue standard + types ajoutés par l'utilisateur."""
    standard = _risk_comp.building_types()
    known    = {b["label"] for b in standard}
    custom   = [c for c in _pdb.list_custom_building_types() if c["label"] not in known]
    return {"building_types": standard + custom}


@app.post("/api/risk-assessment/building-types")
async def ra_add_building_type(request: Request):
    """Ajoute un type de bâtiment personnalisé."""
    body  = await request.json()
    label = (body.get("label") or "").strip()
    if not label:
        raise HTTPException(400, "Le libellé du type de bâtiment est requis.")
    if len(label) > 60:
        raise HTTPException(400, "Le libellé ne doit pas dépasser 60 caractères.")
    _pdb.add_custom_building_type(label, (body.get("icon") or "fa-building").strip())
    return {"label": label, "status": "cree"}


@app.get("/api/risk-assessment/list")
async def ra_list(sector: Optional[str] = None, client: Optional[str] = None):
    """Évaluations existantes, pour reprendre une visite en cours."""
    return {"assessments": _pdb.list_risk_assessments(sector, client)}


def _ra_payload(assessment_id: int) -> dict:
    data = _pdb.get_risk_assessment(assessment_id)
    if not data:
        raise HTTPException(404, "Évaluation introuvable.")
    data["components"] = _risk_comp.components_grouped(data["building_type"])
    return data


@app.post("/api/risk-assessment/start")
async def ra_start(request: Request):
    """Crée une évaluation et ses étages."""
    body          = await request.json()
    sector        = (body.get("sector") or "").strip()
    client        = (body.get("client") or "").strip()
    building_type = (body.get("building_type") or "").strip()

    if not building_type:
        raise HTTPException(400, "Le type de bâtiment est requis.")
    if not client:
        raise HTTPException(400, "Le nom du client est requis.")
    try:
        num_floors = int(body.get("num_floors", 1))
    except (TypeError, ValueError):
        raise HTTPException(400, "Le nombre de niveaux doit être un entier.")
    if not 1 <= num_floors <= 50:
        raise HTTPException(400, "Le nombre de niveaux doit être compris entre 1 et 50.")

    assessment_id = _pdb.create_risk_assessment(
        sector=sector or "non_precise", client=client,
        building_type=building_type, num_floors=num_floors,
    )
    return _ra_payload(assessment_id)


@app.get("/api/risk-assessment/{assessment_id}")
async def ra_get(assessment_id: int):
    """Évaluation, ses étages et les composants proposés pour ce bâtiment."""
    return _ra_payload(assessment_id)


@app.put("/api/risk-assessment/{assessment_id}/floors/{floor_number}")
async def ra_rename_floor(assessment_id: int, floor_number: int, request: Request):
    """Renomme un niveau (« Salle serveur », « Archives », …)."""
    body  = await request.json()
    label = (body.get("floor_label") or "").strip()
    if not label:
        raise HTTPException(400, "Le libellé du niveau est requis.")
    if not _pdb.rename_floor(assessment_id, floor_number, label[:80]):
        raise HTTPException(404, "Niveau introuvable.")
    return {"floor_number": floor_number, "floor_label": label[:80]}


@app.get("/api/risk-assessment/{assessment_id}/floors/{floor_number}/notes")
async def ra_floor_notes(assessment_id: int, floor_number: int):
    """Notes déjà saisies sur un niveau, avec leurs photos."""
    floor = _pdb.get_floor(assessment_id, floor_number)
    if not floor:
        raise HTTPException(404, "Niveau introuvable.")
    notes = _pdb.list_floor_notes(floor["id"])
    for n in notes:
        for p in n["photos"]:
            p["url"]   = f"/api/risk-assessment/{assessment_id}/photos/{p['id']}"
            p["thumb"] = p["url"] + "?vignette=1"
    return {"floor": floor, "notes": notes}


@app.post("/api/risk-assessment/{assessment_id}/notes")
async def ra_create_note(
    assessment_id: int,
    floor_number: int        = Form(...),
    component: str           = Form(...),
    note_text: str           = Form(...),
    condition: str           = Form("non_renseigne"),
    risk_level: str          = Form("non_renseigne"),
    photos: List[UploadFile] = File(default=[]),
):
    """Crée une note sur un composant d'un niveau et y rattache les photos."""
    assessment = _pdb.get_risk_assessment(assessment_id)
    if not assessment:
        raise HTTPException(404, "Évaluation introuvable.")

    # Le niveau est résolu par son numéro AU SEIN de cette évaluation : une note
    # ne peut pas atterrir sur le niveau d'une autre évaluation.
    floor = _pdb.get_floor(assessment_id, floor_number)
    if not floor:
        raise HTTPException(404, "Niveau introuvable pour cette évaluation.")

    component = (component or "").strip()
    note_text = (note_text or "").strip()
    if not component:
        raise HTTPException(400, "Le composant est requis.")
    if not note_text:
        raise HTTPException(400, "L'observation ne peut pas être vide.")
    if condition not in _RA_CONDITIONS:
        condition = "non_renseigne"
    if risk_level not in _RA_RISK_LEVELS:
        risk_level = "non_renseigne"

    note_id = _pdb.create_note(floor["id"], component, note_text,
                               condition, risk_level)

    photo_dir = _ra_photo_dir(assessment)
    thumb_dir = photo_dir / "vignettes"
    saved, rejetees = [], []

    for upload in photos or []:
        if not upload or not upload.filename:
            continue
        raw = await upload.read()
        if not raw:
            continue
        if len(raw) > _RA_PHOTO_MAX_BYTES:
            rejetees.append({"fichier": upload.filename, "motif": "fichier trop volumineux"})
            continue

        try:
            img = Image.open(io.BytesIO(raw))
            img.load()
        except Exception:
            rejetees.append({"fichier": upload.filename, "motif": "format d'image non reconnu"})
            continue

        meta = _ra_exif(img)
        img  = ImageOps.exif_transpose(img)          # redresse la photo
        if img.mode not in ("RGB", "L"):
            img = img.convert("RGB")

        base  = f"p{note_id}_{_uuid_mod.uuid4().hex[:8]}.jpg"
        full  = photo_dir / base
        thumb = thumb_dir / base

        grande = img.copy()
        grande.thumbnail((_RA_MAX_EDGE, _RA_MAX_EDGE), Image.LANCZOS)
        grande.save(full, format="JPEG", quality=85, optimize=True)

        img.thumbnail((_RA_THUMB_EDGE, _RA_THUMB_EDGE), Image.LANCZOS)
        img.save(thumb, format="JPEG", quality=75, optimize=True)

        meta.update({
            "width":     grande.width,
            "height":    grande.height,
            "filesize":  full.stat().st_size,
            "mime_type": "image/jpeg",
        })

        photo_id = _pdb.add_photo_to_note(
            note_id, upload.filename,
            str(full.relative_to(_pdb.PROJECTS_DIR)),
            str(thumb.relative_to(_pdb.PROJECTS_DIR)),
            meta, condition, risk_level,
        )
        saved.append({
            "photo_id": photo_id,
            "url":      f"/api/risk-assessment/{assessment_id}/photos/{photo_id}",
            "thumb":    f"/api/risk-assessment/{assessment_id}/photos/{photo_id}?vignette=1",
        })

    return {
        "note_id":      note_id,
        "floor_number": floor_number,
        "component":    component,
        "photos":       saved,
        "rejetees":     rejetees,
    }


@app.delete("/api/risk-assessment/{assessment_id}/notes/{note_id}")
async def ra_delete_note(assessment_id: int, note_id: int):
    """Supprime une note et ses photos."""
    if not _pdb.delete_note(note_id):
        raise HTTPException(404, "Note introuvable.")
    return {"status": "supprimee"}


@app.get("/api/risk-assessment/{assessment_id}/photos/{photo_id}")
async def ra_photo(assessment_id: int, photo_id: int, vignette: int = 0):
    """Sert une photo (vignette=1 pour la miniature)."""
    photo = _pdb.get_photo(photo_id)
    if not photo:
        raise HTTPException(404, "Photo introuvable.")

    rel  = photo["thumbnail_path"] if vignette else photo["original_path"]
    root = _pdb.PROJECTS_DIR.resolve()
    path = (root / rel).resolve()

    # Le chemin vient de la base, mais on refuse toute sortie du dossier projet.
    if root not in path.parents or not path.is_file():
        raise HTTPException(404, "Fichier photo introuvable.")

    return Response(content=path.read_bytes(), media_type="image/jpeg",
                    headers={"Cache-Control": "private, max-age=3600"})


@app.get("/api/risk-assessment/{assessment_id}/export")
async def ra_export(assessment_id: int):
    """Export structuré de l'évaluation, destiné à l'agent de synthèse."""
    data = _pdb.get_assessment_export(assessment_id)
    if not data:
        raise HTTPException(404, "Évaluation introuvable.")
    return data
