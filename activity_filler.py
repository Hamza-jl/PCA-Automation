"""
activity_filler.py
==================
Fill activities into an already-generated BIA fiche (.docx).

Given a fiche and a list of activities, this module updates four sections:

  Section 3  – "Présentation générale de l'activité"   → name + description
  Section 4  – "Description et criticité des activités" → one row per activity
  Section 5.2 – "Évaluation des impacts"                → one table per activity
  Section 5.3 – "Durée maximale d'interruption"         → one row per activity
"""
from __future__ import annotations

import re
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


# ─── low-level helpers ────────────────────────────────────────────────────────

def _elem_text(elem) -> str:
    """Return all text inside an XML element, joined."""
    return "".join((t.text or "") for t in elem.iter(f"{{{_W}}}t")).strip()


def _cell_text(cell) -> str:
    return cell.text.strip()


def _set_para_text(para_elem, text: str) -> None:
    """Replace all <w:r> children of a paragraph element with one run."""
    for r in list(para_elem.findall(f"{{{_W}}}r")):
        para_elem.remove(r)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    if text and (text[0] == " " or text[-1] == " "):
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    r.append(t)
    para_elem.append(r)


def _clear_data_rows(table, header_rows: int = 1) -> None:
    tbl = table._tbl
    for row in list(tbl.findall(f"{{{_W}}}tr"))[header_rows:]:
        tbl.remove(row)


def _clone_row_with_values(template_tr, values: list[str]):
    """Deep-clone template_tr and fill its cells with values."""
    new_tr = deepcopy(template_tr)
    tcs = new_tr.findall(f"{{{_W}}}tc")
    for tc, val in zip(tcs, values):
        paras = tc.findall(f".//{{{_W}}}p")
        if paras:
            _set_para_text(paras[0], val)
            for extra in paras[1:]:
                tc.remove(extra)
    return new_tr


# ─── public API ───────────────────────────────────────────────────────────────

def fill_activities(
    docx_path: Path,
    activities: list[dict],
    output_path: Path,
) -> None:
    """
    Fill activities into a BIA fiche.

    Parameters
    ----------
    docx_path   : source .docx (generated fiche)
    activities  : list of {'name': str, 'description': str}
    output_path : destination path for the updated .docx
    """
    if not activities:
        return

    doc = Document(str(docx_path))

    _update_presentation(doc, activities)
    _update_description_criticite(doc, activities)
    _update_evaluation_impacts(doc, activities)
    _update_dmia(doc, activities)

    doc.save(str(output_path))


# ─── section updaters ─────────────────────────────────────────────────────────

def _update_presentation(doc: Document, activities: list[dict]) -> None:
    """Section 3 – update the 'Présentation générale de l'activité' cell."""
    for table in doc.tables:
        for row in table.rows:
            if len(row.cells) < 2:
                continue
            label = _cell_text(row.cells[0]).lower()
            if "pr" in label and "sentation" in label and "g" in label and "rale" in label:
                tc = row.cells[1]._tc
                paras = tc.findall(f"{{{_W}}}p")
                lines = (
                    ["Les activités principales de l'entité sont les suivantes :"]
                    + [f"- {a['name']} : {a['description']}" for a in activities]
                )
                if paras:
                    _set_para_text(paras[0], lines[0])
                    for extra in paras[1:]:
                        tc.remove(extra)
                    for line in lines[1:]:
                        new_p = deepcopy(paras[0])
                        _set_para_text(new_p, line)
                        tc.append(new_p)
                return


def _update_description_criticite(doc: Document, activities: list[dict]) -> None:
    """Section 4 – one data row per activity."""
    for table in doc.tables:
        if not table.rows:
            continue
        combined = " ".join(_cell_text(c) for c in table.rows[0].cells).lower()
        if "activit" in combined and "criticit" in combined:
            # Save a template row before clearing
            template_tr = (
                deepcopy(table.rows[1]._tr) if len(table.rows) > 1 else None
            )
            _clear_data_rows(table, header_rows=1)
            if template_tr is not None:
                n_cols = len(table.rows[0].cells)
                for act in activities:
                    vals = [act["name"]] + [""] * (n_cols - 1)
                    table._tbl.append(_clone_row_with_values(template_tr, vals))
            return


def _update_evaluation_impacts(doc: Document, activities: list[dict]) -> None:
    """Section 5.2 – one impact table per activity (duplicate the template table)."""
    body = doc.element.body
    body_children = list(body)

    # Locate the impact EVALUATION table (Section 5.2).
    # Distinguish it from the impact MATRIX (Section 5.1) which has a "Poids" column.
    # The evaluation table has:
    #   row0[0] = activity name placeholder ('xx')   — short, NO "Poids"
    #   row1[0] = 'Image de marque'
    #   typically 3 columns (activity | interruption A | interruption B)
    impact_tbl = None
    impact_idx = None
    for i, child in enumerate(body_children):
        if child.tag != f"{{{_W}}}tbl":
            continue
        rows = child.findall(f"{{{_W}}}tr")
        if len(rows) < 2:
            continue
        row0_tcs = rows[0].findall(f"{{{_W}}}tc")
        row1_tcs = rows[1].findall(f"{{{_W}}}tc")
        if not row0_tcs or not row1_tcs:
            continue
        row1_c0 = _elem_text(row1_tcs[0]).lower()
        row0_all = " ".join(_elem_text(c) for c in row0_tcs).lower()
        # Evaluation table: "image de marque" in row1[0] AND "poids" NOT in row0
        if "image de marque" in row1_c0 and "poids" not in row0_all:
            impact_tbl = child
            impact_idx = i
            break

    if impact_tbl is None:
        return

    # Look for the caption paragraph immediately before the table
    caption_elem = None
    if impact_idx and impact_idx > 0:
        prev = body_children[impact_idx - 1]
        if prev.tag == f"{{{_W}}}p":
            txt = _elem_text(prev).lower()
            if "tableau" in txt or "xx" in txt:
                caption_elem = prev

    # Update first activity in the existing (template) table
    _set_impact_row0_name(impact_tbl, activities[0]["name"])
    if caption_elem is not None:
        _replace_xx_in_para(caption_elem, activities[0]["name"])

    # Duplicate for activities 2..N, inserting after the previous element
    insert_after = impact_tbl
    for act in activities[1:]:
        new_tbl = deepcopy(impact_tbl)
        _set_impact_row0_name(new_tbl, act["name"])

        if caption_elem is not None:
            new_cap = deepcopy(caption_elem)
            _replace_xx_in_para(new_cap, act["name"])
            # Insert caption first, then table after it
            insert_after.addnext(new_tbl)
            insert_after.addnext(new_cap)
        else:
            insert_after.addnext(new_tbl)

        insert_after = new_tbl


def _update_dmia(doc: Document, activities: list[dict]) -> None:
    """Section 5.3 – one row per activity in the DMIA table."""
    for table in doc.tables:
        if not table.rows:
            continue
        combined = " ".join(_cell_text(c) for c in table.rows[0].cells).lower()
        if "processus" in combined and "dmia" in combined:
            template_tr = (
                deepcopy(table.rows[1]._tr) if len(table.rows) > 1 else None
            )
            _clear_data_rows(table, header_rows=1)
            if template_tr is not None:
                n_cols = len(table.rows[0].cells)
                for act in activities:
                    vals = [act["name"]] + [""] * (n_cols - 1)
                    table._tbl.append(_clone_row_with_values(template_tr, vals))
            return


# ─── XML helpers ──────────────────────────────────────────────────────────────

def _set_impact_row0_name(tbl_elem, name: str) -> None:
    """Set activity name in row 0, cell 0 of an impact table XML element."""
    rows = tbl_elem.findall(f"{{{_W}}}tr")
    if not rows:
        return
    cells = rows[0].findall(f"{{{_W}}}tc")
    if not cells:
        return
    paras = cells[0].findall(f".//{{{_W}}}p")
    if paras:
        _set_para_text(paras[0], name)


def _replace_xx_in_para(para_elem, replacement: str) -> None:
    """Replace the first occurrence of 'xx' in a paragraph with replacement."""
    for t in para_elem.iter(f"{{{_W}}}t"):
        if t.text and re.search(r"\bxx\b", t.text, re.IGNORECASE):
            t.text = re.sub(r"\bxx\b", replacement, t.text, count=1, flags=re.IGNORECASE)
            return
