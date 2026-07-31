"""
questionnaires_db.py — Collecte de données module (Gestion des risques).

Two sub-features share this module:
  - Questionnaires spécifiques: one global reference (template) questionnaire
    per type (Bâtiment - Siège, Bâtiment - Agence, RH, Fournisseurs, Cyber,
    Gouvernance de crise, Télétravail), fully editable — rows AND columns.
  - Réponses: a client's filled-in response file, compared against the
    reference to find unanswered / missing questions, then optionally saved
    to the DB under a sector/client.

Columns are dynamic: whatever header row the Excel file has becomes the
column list verbatim (no fixed field names), so users can add, delete, or
rename columns for any of the 7 types without a code change. Each row is a
plain {column_name: value} dict.

Unlike risk_analysis.py's in-memory-only Excel cache, the original file bytes
are persisted to disk under projects_data/ so saved questionnaires and
template files both survive a server restart.
"""
from __future__ import annotations

import json
import re
import sqlite3
import unicodedata
from copy import copy as _copy_style
from io import BytesIO
from pathlib import Path
from typing import Optional

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

DB_PATH       = Path(__file__).parent / "projects.db"
BASE_DIR      = Path(__file__).parent
FILES_DIR     = BASE_DIR / "projects_data"
TEMPLATES_DIR = FILES_DIR / "_templates"

QUESTIONNAIRE_TYPES = [
    {"slug": "batiment_siege",    "label": "Bâtiment - Siège",     "icon": "fa-solid fa-building"},
    {"slug": "batiment_agence",   "label": "Bâtiment - Agence",    "icon": "fa-solid fa-store"},
    {"slug": "rh",                "label": "RH",                   "icon": "fa-solid fa-users"},
    {"slug": "fournisseurs",      "label": "Fournisseurs",         "icon": "fa-solid fa-truck-field"},
    {"slug": "cyber",             "label": "Cyber",                "icon": "fa-solid fa-lock"},
    {"slug": "gouvernance_crise", "label": "Gouvernance de crise", "icon": "fa-solid fa-triangle-exclamation"},
    {"slug": "teletravail",       "label": "Télétravail",          "icon": "fa-solid fa-house-laptop"},
]
_TYPE_SLUGS = {t["slug"] for t in QUESTIONNAIRE_TYPES}

_RESPONSE_OPTIONS = ["À renseigner", "Oui", "Non", "N.A."]


def _normalize(s: str) -> str:
    return unicodedata.normalize("NFKD", (s or "").lower()).encode("ascii", "ignore").decode().strip()


def _slug(s: str) -> str:
    return re.sub(r'[^\w\-]', '_', (s or '').strip())[:60]


def find_column(columns: list[str], name: str) -> Optional[str]:
    """Case/accent-insensitive lookup of a column by name (e.g. 'reponse' -> 'Réponse')."""
    target = _normalize(name)
    for c in columns:
        if _normalize(c) == target:
            return c
    return None


# ─────────────────────────────────────────────────────────────────────────────
# Header detection — fully generic: whatever non-empty cells sit on the first
# row (scanned 1-15) with >= 4 of them become the columns, verbatim. This is
# what makes columns dynamic per questionnaire type instead of hardcoded.
# ─────────────────────────────────────────────────────────────────────────────

def _detect_header(ws) -> tuple[int, int, list[tuple[int, str]]]:
    """
    Returns (header_row, data_start_row, col_defs) where col_defs is a list of
    (1-based column index, column name) in left-to-right order, taken directly
    from whatever text is in the Excel header cells.
    """
    for row in range(1, 16):
        cells: list[tuple[int, str]] = []
        for col in range(1, 30):
            v = ws.cell(row=row, column=col).value
            if v not in (None, ""):
                cells.append((col, str(v).strip()))
        if len(cells) >= 4:
            return row, row + 1, cells

    # Fallback: matches the layout of the example file (header row 5)
    header_row = 5
    col_defs = [
        (2, "I.D"), (3, "Axe"), (4, "Question"), (5, "Owner"),
        (6, "Réponse"), (7, "Commentaires"), (8, "Moyens de contournement mis en place"), (9, "Preuves"),
    ]
    return header_row, header_row + 1, col_defs


def _cell(ws, row: int, col: Optional[int]):
    if not col:
        return ""
    v = ws.cell(row=row, column=col).value
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        return int(v)
    return v


# ─────────────────────────────────────────────────────────────────────────────
# Excel parsing
# ─────────────────────────────────────────────────────────────────────────────

def parse_questionnaire_excel(
    file_bytes: bytes, sheet_name: str | None = None
) -> tuple[list[str], list[dict], str, list[str]]:
    """Parse a questionnaire Excel file. Returns (columns, rows, detected_sheet_name, available_sheets)."""
    wb = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True)
    sheets = wb.sheetnames

    if sheet_name is None:
        for s in sheets:
            nl = _normalize(s)
            if "prealable" in nl or "questionnaire" in nl:
                sheet_name = s
                break
        if sheet_name is None:
            sheet_name = sheets[0]

    ws = wb[sheet_name]
    _hrow, data_start, col_defs = _detect_header(ws)
    columns = [name for _, name in col_defs]

    rows: list[dict] = []
    for r in range(data_start, (ws.max_row or data_start) + 1):
        row_vals: dict = {}
        any_val = False
        for col_idx, name in col_defs:
            v = _cell(ws, r, col_idx)
            if v != "":
                any_val = True
            row_vals[name] = str(v) if v != "" else ""
        if any_val:
            rows.append(row_vals)

    return columns, rows, sheet_name, sheets


# ─────────────────────────────────────────────────────────────────────────────
# Excel export — full rewrite (not raw XML patch) so rows AND columns can be
# added/removed/renamed since the reference questionnaire was first uploaded.
# ─────────────────────────────────────────────────────────────────────────────

def export_questionnaire_excel(
    original_bytes: bytes, columns: list[str], rows: list[dict], sheet_name: str
) -> bytes:
    wb = openpyxl.load_workbook(BytesIO(original_bytes))
    ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.worksheets[0]
    orig_header_row, orig_data_start, orig_col_defs = _detect_header(ws)
    start_col = orig_col_defs[0][0] if orig_col_defs else 2

    def _style_of(cell) -> dict:
        return {
            "font": _copy_style(cell.font), "border": _copy_style(cell.border),
            "fill": _copy_style(cell.fill), "number_format": cell.number_format,
            "protection": _copy_style(cell.protection), "alignment": _copy_style(cell.alignment),
        }

    def _apply_style(cell, st: Optional[dict]) -> None:
        if not st:
            return
        cell.font = st["font"]; cell.border = st["border"]; cell.fill = st["fill"]
        cell.number_format = st["number_format"]; cell.protection = st["protection"]
        cell.alignment = st["alignment"]

    header_style = _style_of(ws.cell(row=orig_header_row, column=start_col))
    data_style = (
        _style_of(ws.cell(row=orig_data_start, column=start_col))
        if ws.max_row >= orig_data_start else header_style
    )

    # Existing data validations reference ranges that won't line up with the
    # rewritten table — drop them, we rebuild the Réponse dropdown below.
    ws.data_validations.dataValidation = []

    # Clear a rectangle big enough to cover both the old and new table extents
    old_max_row = ws.max_row
    old_col_span = len(orig_col_defs)
    new_col_span = len(columns)
    last_row = max(old_max_row, orig_data_start + len(rows) - 1, orig_header_row)
    last_col = start_col + max(old_col_span, new_col_span) - 1
    for r in range(orig_header_row, last_row + 1):
        for c in range(start_col, last_col + 1):
            ws.cell(row=r, column=c).value = None

    # Write header
    for i, col_name in enumerate(columns):
        cell = ws.cell(row=orig_header_row, column=start_col + i)
        cell.value = col_name
        _apply_style(cell, header_style)

    # Write rows
    for ri, row in enumerate(rows):
        rr = orig_data_start + ri
        for i, col_name in enumerate(columns):
            cell = ws.cell(row=rr, column=start_col + i)
            cell.value = row.get(col_name, "")
            _apply_style(cell, data_style)

    # Re-apply the dropdown to whichever column is currently named "Réponse"
    resp_name = find_column(columns, "reponse")
    if resp_name and rows:
        resp_idx = columns.index(resp_name)
        col_letter = get_column_letter(start_col + resp_idx)
        dv = DataValidation(
            type="list",
            formula1='"{}"'.format(",".join(_RESPONSE_OPTIONS)),
            allow_blank=True,
        )
        dv.add(f"{col_letter}{orig_data_start}:{col_letter}{orig_data_start + len(rows) - 1}")
        ws.add_data_validation(dv)

    out = BytesIO()
    wb.save(out)
    return out.getvalue()


# ─────────────────────────────────────────────────────────────────────────────
# DB
# ─────────────────────────────────────────────────────────────────────────────

def _conn() -> sqlite3.Connection:
    con = sqlite3.connect(str(DB_PATH))
    con.row_factory = sqlite3.Row
    return con


def _init_db() -> None:
    FILES_DIR.mkdir(exist_ok=True)
    TEMPLATES_DIR.mkdir(parents=True, exist_ok=True)
    with _conn() as con:
        con.executescript("""
            CREATE TABLE IF NOT EXISTS questionnaires (
                id           INTEGER PRIMARY KEY AUTOINCREMENT,
                sector       TEXT NOT NULL DEFAULT '',
                client       TEXT NOT NULL DEFAULT '',
                qtype        TEXT NOT NULL,
                name         TEXT NOT NULL DEFAULT '',
                is_template  INTEGER NOT NULL DEFAULT 0,
                sheet_name   TEXT DEFAULT '',
                rows_json    TEXT NOT NULL,
                file_path    TEXT DEFAULT '',
                created_at   TEXT DEFAULT (datetime('now')),
                updated_at   TEXT DEFAULT (datetime('now'))
            );
            CREATE INDEX IF NOT EXISTS idx_q_type   ON questionnaires(qtype);
            CREATE INDEX IF NOT EXISTS idx_q_client  ON questionnaires(sector, client);
        """)


def _store_path(sector: str, client: str, qtype: str, is_template: bool, qid: int) -> Path:
    if is_template:
        d = TEMPLATES_DIR / qtype
    else:
        d = FILES_DIR / _slug(sector) / _slug(client) / "questionnaires" / qtype
    d.mkdir(parents=True, exist_ok=True)
    return d / f"{qid}.xlsx"


def _pack(columns: list[str], rows: list[dict]) -> str:
    return json.dumps({"columns": columns, "rows": rows}, ensure_ascii=False)


def _unpack(blob: str) -> tuple[list[str], list[dict]]:
    data = json.loads(blob)
    return data.get("columns", []), data.get("rows", [])


def save_questionnaire(
    name: str, sector: str, client: str, qtype: str,
    columns: list[str], rows: list[dict], sheet: str, original_bytes: bytes,
    is_template: bool = False,
) -> int:
    _init_db()
    if qtype not in _TYPE_SLUGS:
        raise ValueError(f"Type de questionnaire inconnu : {qtype}")

    if is_template:
        sector, client = "", ""

    with _conn() as con:
        if is_template:
            con.execute("UPDATE questionnaires SET is_template=0 WHERE qtype=? AND is_template=1", (qtype,))
        cur = con.execute(
            "INSERT INTO questionnaires (sector, client, qtype, name, is_template, sheet_name, rows_json) "
            "VALUES (?,?,?,?,?,?,?)",
            (sector.strip(), client.strip(), qtype, name.strip(), int(is_template), sheet,
             _pack(columns, rows)),
        )
        qid = cur.lastrowid
        path = _store_path(sector, client, qtype, is_template, qid)
        path.write_bytes(original_bytes)
        con.execute("UPDATE questionnaires SET file_path=? WHERE id=?", (str(path), qid))
    return qid


def update_questionnaire(
    qid: int, columns: list[str], rows: list[dict], name: Optional[str] = None
) -> None:
    _init_db()
    with _conn() as con:
        if name and name.strip():
            con.execute(
                "UPDATE questionnaires SET rows_json=?, name=?, updated_at=datetime('now') WHERE id=?",
                (_pack(columns, rows), name.strip(), qid),
            )
        else:
            con.execute(
                "UPDATE questionnaires SET rows_json=?, updated_at=datetime('now') WHERE id=?",
                (_pack(columns, rows), qid),
            )


def list_questionnaires(sector: str = "", client: str = "", qtype: str = "") -> list[dict]:
    _init_db()
    q = ("SELECT id, sector, client, qtype, name, is_template, created_at, updated_at "
         "FROM questionnaires WHERE is_template=0")
    params: list = []
    if sector:
        q += " AND sector=? COLLATE NOCASE"; params.append(sector)
    if client:
        q += " AND client=? COLLATE NOCASE"; params.append(client)
    if qtype:
        q += " AND qtype=?"; params.append(qtype)
    q += " ORDER BY updated_at DESC"
    with _conn() as con:
        rows = con.execute(q, params).fetchall()
    return [dict(r) for r in rows]


def count_by_type(sector: str = "", client: str = "") -> dict:
    """{qtype: count} of saved (non-template) questionnaires for a client."""
    _init_db()
    q = "SELECT qtype, COUNT(*) as c FROM questionnaires WHERE is_template=0"
    params: list = []
    if sector:
        q += " AND sector=? COLLATE NOCASE"; params.append(sector)
    if client:
        q += " AND client=? COLLATE NOCASE"; params.append(client)
    q += " GROUP BY qtype"
    with _conn() as con:
        rows = con.execute(q, params).fetchall()
    return {r["qtype"]: r["c"] for r in rows}


def load_questionnaire(qid: int) -> Optional[dict]:
    _init_db()
    with _conn() as con:
        row = con.execute("SELECT * FROM questionnaires WHERE id=?", (qid,)).fetchone()
    if not row:
        return None
    d = dict(row)
    d["columns"], d["rows"] = _unpack(d.pop("rows_json"))
    return d


def get_original_bytes(qid: int) -> Optional[bytes]:
    d = load_questionnaire(qid)
    if not d or not d.get("file_path"):
        return None
    p = Path(d["file_path"])
    if not p.exists():
        return None
    return p.read_bytes()


def delete_questionnaire(qid: int) -> None:
    _init_db()
    with _conn() as con:
        row = con.execute("SELECT file_path FROM questionnaires WHERE id=?", (qid,)).fetchone()
        con.execute("DELETE FROM questionnaires WHERE id=?", (qid,))
    if row and row["file_path"]:
        try:
            Path(row["file_path"]).unlink(missing_ok=True)
        except Exception:
            pass


def get_template(qtype: str) -> Optional[dict]:
    _init_db()
    with _conn() as con:
        row = con.execute(
            "SELECT * FROM questionnaires WHERE qtype=? AND is_template=1", (qtype,)
        ).fetchone()
    if not row:
        return None
    d = dict(row)
    d["columns"], d["rows"] = _unpack(d.pop("rows_json"))
    return d


def set_as_template(qid: int) -> int:
    """
    Clone questionnaire `qid` into a new template row for its qtype (leaving
    the original filled instance untouched), unflagging any previous template
    of the same type. Returns the new template's id.
    """
    _init_db()
    with _conn() as con:
        row = con.execute("SELECT * FROM questionnaires WHERE id=?", (qid,)).fetchone()
        if not row:
            raise ValueError("Questionnaire introuvable")
        qtype = row["qtype"]
        con.execute("UPDATE questionnaires SET is_template=0 WHERE qtype=? AND is_template=1", (qtype,))
        cur = con.execute(
            "INSERT INTO questionnaires (sector, client, qtype, name, is_template, sheet_name, rows_json) "
            "VALUES ('', '', ?, ?, 1, ?, ?)",
            (qtype, row["name"], row["sheet_name"], row["rows_json"]),
        )
        new_id = cur.lastrowid
        src = Path(row["file_path"])
        dest = _store_path("", "", qtype, True, new_id)
        dest.write_bytes(src.read_bytes())
        con.execute("UPDATE questionnaires SET file_path=? WHERE id=?", (str(dest), new_id))
    return new_id


def create_from_template(
    qtype: str, sector: str, client: str, name: str, columns: list[str], rows: list[dict]
) -> int:
    """Start a new client questionnaire from the qtype's blank template."""
    tmpl = get_template(qtype)
    if not tmpl:
        raise ValueError("Aucun modèle enregistré pour ce type de questionnaire")
    original_bytes = Path(tmpl["file_path"]).read_bytes()
    return save_questionnaire(
        name, sector, client, qtype, columns, rows, tmpl["sheet_name"], original_bytes, is_template=False
    )


def _make_xlsx(columns: list[str], rows: list[dict]) -> bytes:
    """Generate a minimal xlsx from column/row data (used when no original file exists)."""
    import openpyxl as _xl
    wb = _xl.Workbook()
    ws = wb.active
    for i, col in enumerate(columns, 1):
        ws.cell(row=1, column=i).value = col
    for ri, row_data in enumerate(rows, 2):
        for ci, col in enumerate(columns, 1):
            ws.cell(row=ri, column=ci).value = row_data.get(col, "")
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def save_template_direct(
    qtype: str, columns: list[str], rows: list[dict],
    name: str = "Questionnaire de référence",
) -> int:
    """Save (or update) the reference template for a type — no prior file required."""
    _init_db()
    if qtype not in _TYPE_SLUGS:
        raise ValueError(f"Type de questionnaire inconnu : {qtype}")
    existing = get_template(qtype)
    if existing:
        update_questionnaire(existing["id"], columns, rows, name)
        return existing["id"]
    file_bytes = _make_xlsx(columns, rows)
    return save_questionnaire(name, "", "", qtype, columns, rows, "Sheet", file_bytes, is_template=True)


def save_questionnaire_direct(
    qtype: str, sector: str, client: str, name: str,
    columns: list[str], rows: list[dict],
) -> int:
    """Save a client questionnaire without requiring a pre-existing template."""
    _init_db()
    if qtype not in _TYPE_SLUGS:
        raise ValueError(f"Type de questionnaire inconnu : {qtype}")
    tmpl = get_template(qtype)
    if tmpl and Path(tmpl["file_path"]).exists():
        file_bytes = Path(tmpl["file_path"]).read_bytes()
        sheet = tmpl.get("sheet_name") or "Sheet"
    else:
        file_bytes = _make_xlsx(columns, rows)
        sheet = "Sheet"
    return save_questionnaire(name, sector, client, qtype, columns, rows, sheet, file_bytes, is_template=False)
